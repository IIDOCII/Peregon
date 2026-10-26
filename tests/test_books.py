"""Этап 4: книги — EPUB, FB2, MOBI → EPUB, FB2, PDF, TXT (П18) и безопасность.

Книги генерируются в тестах: FB2 — шаблоном, EPUB 2/3 — zipfile вручную,
MOBI — по спецификации PalmDB/PalmDOC/MOBI/EXTH (без сжатия и с LZ77).
Большинство тестов вызывают задачу ребёнка в процессе тестов (`inproc`);
несколько — через настоящий дочерний процесс.
"""

from __future__ import annotations

import base64
import os
import struct
import threading
import time
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from lxml import etree

from docs_helpers import (
    EVIL_HTML,
    LAUGHS,
    assert_clean,
    bomb_zip,
    file_uri,
    jpeg_bytes,
    lie_about_size,
    make_inproc,
    many_entries_zip,
    no_children,
    png_bytes,
    run,
    secret,  # noqa: F401 — фикстура
    target,
    wait_jobs,
    zip_text,
)
from konverter.converters import _docjob, _isolated, books
from konverter.converters.base import (
    REGISTRY,
    CancelToken,
    Cancelled,
    Category,
    ConversionError,
    Preset,
)
from konverter.core import naming
from konverter.core.grouping import Group, common_targets
from konverter.core.jobs import JobManager, JobStatus
from konverter.core.detect import detect

CONV = books.BookConverter()
OPF = "{http://www.idpf.org/2007/opf}"
DC = "{http://purl.org/dc/elements/1.1/}"
FB = "{http://www.gribuser.ru/xml/fictionbook/2.0}"


@pytest.fixture
def inproc(monkeypatch):
    return make_inproc(monkeypatch)


# --- Генераторы книг ------------------------------------------------------------------


def make_fb2(path: Path, *, title="Война и мир", author=("Лев", "Толстой"), chapters=None,
             cover=True, image=True, extra_body="", doctype="") -> Path:
    chapters = chapters or [
        ("Глава первая", ["Ну, князь, Генуя и Лукка стали не больше как поместья.",
                          "Второй <strong>жирный</strong> и <emphasis>курсив</emphasis> абзац."]),
        ("Глава вторая", ["Анна Павловна кашляла уже несколько дней."]),
        ("Глава третья", ["Последний абзац с ссылкой "
                          '<a l:href="https://ru.wikipedia.org/">вики</a>.']),
    ]
    sections = []
    for index, (head, paras) in enumerate(chapters):
        body = "".join(f"<p>{p}</p>" for p in paras)
        if image and index == 1:
            body += '<image l:href="#pic.png"/>'
        sections.append(f"<section><title><p>{head}</p></title>{body}</section>")
    binaries = ""
    coverpage = ""
    if cover:
        coverpage = '<coverpage><image l:href="#cover.jpg"/></coverpage>'
        binaries += ('<binary id="cover.jpg" content-type="image/jpeg">'
                     + base64.b64encode(jpeg_bytes()).decode() + "</binary>")
    if image:
        binaries += ('<binary id="pic.png" content-type="image/png">'
                     + base64.b64encode(png_bytes()).decode() + "</binary>")
    text = (
        '<?xml version="1.0" encoding="utf-8"?>' + doctype
        + '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" '
        'xmlns:l="http://www.w3.org/1999/xlink"><description><title-info>'
        f"<genre>prose</genre><author><first-name>{author[0]}</first-name>"
        f"<last-name>{author[1]}</last-name></author><book-title>{title}</book-title>"
        f"{coverpage}<lang>ru</lang></title-info></description>"
        f"<body>{''.join(sections)}{extra_body}</body>{binaries}</FictionBook>"
    )
    path.write_bytes(text.encode("utf-8"))
    return path


_XHTML = (
    '<?xml version="1.0" encoding="utf-8"?>{doctype}<html xmlns="http://www.w3.org/1999/xhtml">'
    "<head><title>{title}</title></head><body>{body}</body></html>"
)


def make_epub(path: Path, *, version="2.0", chapters=None, cover=True, extra=None,
              opf_doctype="", chapter_doctype="", nav=True) -> Path:
    """EPUB 2 (NCX) или 3 (nav) с обложкой, картинкой в главе и главами по spine."""
    chapters = chapters or [
        ("Начало", "<h1>Начало</h1><p>Первая глава книги.</p>"
                   '<p><img src="../images/pic.png" alt="рисунок"/></p>'),
        ("Середина", "<h2>Середина</h2><p>Вторая <em>глава</em> книги.</p>"),
        ("Конец", "<h1>Конец</h1><p>Третья глава и <strong>финал</strong>.</p>"),
    ]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        info = zipfile.ZipInfo("mimetype")
        info.compress_type = zipfile.ZIP_STORED
        archive.writestr(info, b"application/epub+zip")
        archive.writestr("META-INF/container.xml", (
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="OPS/book.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>"))
        manifest = []
        spine = []
        # Порядок в spine — обратный порядку в manifest: проверяем, что читается spine.
        for index, (title, body) in enumerate(chapters):
            name = f"text/c{index}.xhtml"
            archive.writestr("OPS/" + name, _XHTML.format(doctype=chapter_doctype, title=title,
                                                           body=body))
            manifest.insert(0, f'<item id="c{index}" href="{name}" '
                               'media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="c{index}"/>')
        archive.writestr("OPS/images/pic.png", png_bytes())
        manifest.append('<item id="pic" href="images/pic.png" media-type="image/png"/>')
        meta = ""
        if cover:
            archive.writestr("OPS/images/cover.jpg", jpeg_bytes())
            props = ' properties="cover-image"' if version.startswith("3") else ""
            manifest.append(f'<item id="cov" href="images/cover.jpg" media-type="image/jpeg"{props}/>')
            if not version.startswith("3"):
                meta = '<meta name="cover" content="cov"/>'
            archive.writestr("OPS/text/cover.xhtml", _XHTML.format(
                doctype="", title="Обложка", body='<div><img src="../images/cover.jpg"/></div>'))
            manifest.append('<item id="coverpage" href="text/cover.xhtml" '
                            'media-type="application/xhtml+xml"/>')
            spine.insert(0, '<itemref idref="coverpage"/>')
        points = "".join(
            f'<navPoint id="n{i}" playOrder="{i + 1}"><navLabel><text>Глава: {t}</text>'
            f'</navLabel><content src="text/c{i}.xhtml"/></navPoint>'
            for i, (t, _b) in enumerate(chapters))
        archive.writestr("OPS/toc.ncx", (
            '<?xml version="1.0"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" '
            f'version="2005-1"><head/><docTitle><text>Книга</text></docTitle>'
            f"<navMap>{points}</navMap></ncx>"))
        manifest.append('<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>')
        if version.startswith("3") and nav:
            links = "".join(f'<li><a href="text/c{i}.xhtml">Навигация {t}</a></li>'
                            for i, (t, _b) in enumerate(chapters))
            archive.writestr("OPS/nav.xhtml", (
                '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" '
                'xmlns:epub="http://www.idpf.org/2007/ops"><head><title>nav</title></head><body>'
                f'<nav epub:type="toc"><ol>{links}</ol></nav></body></html>'))
            manifest.append('<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" '
                            'properties="nav"/>')
        for name, data in (extra or {}).items():
            archive.writestr(name, data)
        archive.writestr("OPS/book.opf", (
            f'<?xml version="1.0" encoding="utf-8"?>{opf_doctype}'
            f'<package xmlns="http://www.idpf.org/2007/opf" version="{version}" '
            'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            '<dc:identifier id="id">urn:uuid:1</dc:identifier><dc:title>Тестовая книга</dc:title>'
            "<dc:creator>Фёдор Достоевский</dc:creator><dc:language>ru</dc:language>"
            f"{meta}</metadata><manifest>{''.join(manifest)}</manifest>"
            f'<spine toc="ncx">{"".join(spine)}</spine></package>'))
    return path


# --- MOBI (PalmDB + PalmDOC + MOBI + EXTH) -------------------------------------------


def lz77(data: bytes) -> bytes:
    """Простой упаковщик PalmDOC (литералы, «пробел+символ», обратные ссылки)."""
    out = bytearray()
    i = 0
    n = len(data)
    while i < n:
        best = 0
        dist = 0
        for length in range(10, 2, -1):
            if i + length > n:
                continue
            pos = data.rfind(data[i:i + length], max(0, i - 2047), i)
            if pos >= 0 and i - pos <= 2047:
                best, dist = length, i - pos
                break
        if best:
            code = 0x8000 | (dist << 3) | (best - 3)
            out += struct.pack(">H", code)
            i += best
            continue
        c = data[i]
        if c == 0x20 and i + 1 < n and 0x40 <= data[i + 1] <= 0x7F:
            out.append(data[i + 1] ^ 0x80)
            i += 2
            continue
        if c == 0 or 0x09 <= c <= 0x7F:
            out.append(c)
            i += 1
            continue
        j = i
        while j < n and j - i < 8 and not (data[j] == 0 or 0x09 <= data[j] <= 0x7F):
            j += 1
        out.append(j - i)
        out += data[i:j]
        i = j
    return bytes(out)


MOBI_HTML = (
    "<html><head><guide></guide></head><body>"
    "<h1>Первая глава</h1><p>Текст <b>первой</b> главы, ёжик.</p>"
    '<p><img recindex="00002"/></p>'
    "<mbp:pagebreak/>"
    "<h1>Вторая глава</h1><p>Текст второй главы.</p><script>alert(1)</script>"
    "<mbp:pagebreak/>"
    "<h1>Третья глава</h1><p>Конец книги.</p>"
    "</body></html>"
)


def make_mobi(path: Path, html: str = MOBI_HTML, *, compression=2, encryption=0, version=6,
              title="Мобильная книга", author="Антон Чехов", extra_flags=0,
              with_images=True) -> Path:
    text = html.encode("utf-8")
    chunks = [text[i:i + 4096] for i in range(0, len(text), 4096)]
    text_records = []
    for chunk in chunks:
        body = lz77(chunk) if compression == 2 else chunk
        if extra_flags & 1:
            body += b"\x00"  # многобайтный хвост: 0 лишних байт + байт счётчика
        if extra_flags & 2:
            body += b"ZZ" + bytes([0x80 | 3])  # хвостовая запись из 3 байт
        text_records.append(body)
    images = [jpeg_bytes(), png_bytes()] if with_images else []
    first_image = 1 + len(text_records) if images else 0xFFFFFFFF
    exth_items = [(100, author.encode("utf-8")), (503, title.encode("utf-8")),
                  (524, b"ru")]
    if images:
        exth_items.append((201, struct.pack(">I", 0)))  # обложка — первая картинка
    exth_body = b"".join(struct.pack(">II", k, len(v) + 8) + v for k, v in exth_items)
    exth = b"EXTH" + struct.pack(">II", 12 + len(exth_body), len(exth_items)) + exth_body
    exth += b"\0" * ((4 - len(exth) % 4) % 4)
    header_len = 0xE8
    rec0 = bytearray(16 + header_len)
    struct.pack_into(">HHIHHHH", rec0, 0, compression, 0, len(text), len(text_records), 4096,
                     encryption, 0)
    rec0[16:20] = b"MOBI"
    struct.pack_into(">IIII", rec0, 20, header_len, 2, 65001, 12345)
    struct.pack_into(">I", rec0, 36, version)
    full_name = "Полное имя".encode("utf-8")
    name_off = len(rec0) + len(exth)
    struct.pack_into(">II", rec0, 0x54, name_off, len(full_name))
    struct.pack_into(">I", rec0, 0x6C, first_image)
    struct.pack_into(">I", rec0, 0x80, 0x40)
    struct.pack_into(">H", rec0, 0xF2, extra_flags)
    rec0 = bytes(rec0) + exth + full_name + b"\0\0"
    records = [rec0] + text_records + images + [b"FLIS\0\0\0\x08", b"\xe9\x8e\r\n"]
    count = len(records)
    header = bytearray(78)
    header[0:6] = b"mobi\0\0"
    header[60:68] = b"BOOKMOBI"
    struct.pack_into(">H", header, 76, count)
    offset = 78 + 8 * count + 2
    table = bytearray()
    for index, record in enumerate(records):
        table += struct.pack(">I", offset) + bytes([0]) + index.to_bytes(3, "big")
        offset += len(record)
    path.write_bytes(bytes(header) + bytes(table) + b"\0\0" + b"".join(records))
    return path


# --- Разбор результатов --------------------------------------------------------------


def epub_info(path: Path) -> dict:
    """Проверяет структуру EPUB и возвращает метаданные, главы (тексты) и картинки."""
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        assert infos[0].filename == "mimetype"
        assert infos[0].compress_type == zipfile.ZIP_STORED
        assert archive.read("mimetype") == b"application/epub+zip"
        container = etree.fromstring(archive.read("META-INF/container.xml"))
        opf_path = container.find(".//{*}rootfile").get("full-path")
        opf = etree.fromstring(archive.read(opf_path))
        base = os.path.dirname(opf_path)
        assert opf.get("version") == "3.0"
        title = opf.find(f".//{DC}title").text
        authors = [e.text for e in opf.iter(f"{DC}creator")]
        assert opf.find(f".//{OPF}meta[@property='dcterms:modified']") is not None
        items = {i.get("id"): i for i in opf.iter(f"{OPF}item")}
        nav = [i for i in items.values() if "nav" in (i.get("properties") or "")]
        assert len(nav) == 1
        spine = opf.find(f"{OPF}spine")
        assert spine.get("toc") in items
        ncx = etree.fromstring(archive.read(base + "/" + items[spine.get("toc")].get("href")))
        nav_labels = [t.text for t in ncx.iter("{http://www.daisy.org/z3986/2005/ncx/}text")]
        chapters = []
        for ref in spine.iter(f"{OPF}itemref"):
            href = items[ref.get("idref")].get("href")
            root = etree.fromstring(archive.read(base + "/" + href))
            chapters.append(" ".join(" ".join(root.find("{*}body").itertext()).split()))
        cover = [i for i in items.values() if "cover-image" in (i.get("properties") or "")]
        images = [i.get("href") for i in items.values()
                  if (i.get("media-type") or "").startswith("image/")]
        for href in images:
            archive.read(base + "/" + href)
    return {"title": title, "authors": authors, "chapters": chapters, "cover": cover,
            "images": images, "ncx": nav_labels}


def fb2_info(path: Path) -> dict:
    root = etree.fromstring(path.read_bytes())
    info = root.find(f"{FB}description/{FB}title-info")
    author = info.find(f"{FB}author")
    name = " ".join(e.text for e in author if e.text)
    sections = [" ".join(" ".join(s.itertext()).split()) for s in root.find(f"{FB}body")
                if s.tag == f"{FB}section"]
    binaries = {b.get("id"): base64.b64decode(b.text) for b in root.iter(f"{FB}binary")}
    cover = info.find(f"{FB}coverpage/{FB}image")
    return {"title": info.find(f"{FB}book-title").text, "author": name, "sections": sections,
            "binaries": binaries,
            "cover": cover.get("{http://www.w3.org/1999/xlink}href") if cover is not None else None,
            "lang": info.find(f"{FB}lang").text}


def pdf_text(path: Path) -> str:
    doc = pdfium.PdfDocument(str(path))
    try:
        parts = []
        for index in range(len(doc)):
            page = doc[index]
            textpage = page.get_textpage()
            parts.append(textpage.get_text_range())
            textpage.close()
            page.close()
        return "\n".join(parts)
    finally:
        doc.close()


def _order(text: str, *parts: str) -> None:
    positions = [text.find(p) for p in parts]
    assert all(p >= 0 for p in positions), (parts, positions)
    assert positions == sorted(positions), (parts, positions)


# --- Варианты ---------------------------------------------------------------------------


def test_targets_order_and_registry():
    assert isinstance(REGISTRY[Category.BOOK], books.BookConverter)
    assert [t.id for t in CONV.targets("epub")] == ["fb2", "pdf", "txt"]
    assert [t.id for t in CONV.targets("fb2")] == ["epub", "pdf", "txt"]
    assert [t.id for t in CONV.targets("mobi")] == ["epub", "fb2", "pdf", "txt"]
    assert CONV.targets("docx") == []
    for t in CONV.targets("mobi"):
        assert t.is_long and not t.many_to_one
        assert t.uses_preset == (t.id == "pdf")
        assert t.label and not t.label.startswith("target_")
    assert CONV.queue == "docs"


def test_group_of_mixed_books_offers_all(tmp_path):
    from konverter.converters.base import DetectedFormat

    files = [(tmp_path / "a.epub", DetectedFormat("epub", Category.BOOK, "EPUB")),
             (tmp_path / "b.fb2", DetectedFormat("fb2", Category.BOOK, "FB2"))]
    ids = [t.id for t in common_targets(Group(Category.BOOK, files))]
    assert ids == ["epub", "fb2", "pdf", "txt"]


# --- П18: EPUB ↔ FB2 ---------------------------------------------------------------------


def test_p18_fb2_to_epub_and_back(tmp_path, inproc):
    src = make_fb2(tmp_path / "книга.fb2")
    assert detect(src).fmt == "fb2"
    epub = run(CONV, src, tmp_path / "out.epub", "fb2", "epub")
    assert detect(epub).fmt == "epub"
    info = epub_info(epub)
    assert info["title"] == "Война и мир"
    assert info["authors"] == ["Лев Толстой"]
    assert len(info["cover"]) == 1
    assert len(info["images"]) == 2  # обложка и картинка главы
    text = " | ".join(info["chapters"])
    _order(text, "Глава первая", "Генуя и Лукка", "Глава вторая", "Анна Павловна",
           "Глава третья", "вики")
    assert info["ncx"][1:] == ["Глава первая", "Глава вторая", "Глава третья"]
    # Жирный/курсив и ссылка — в XHTML главы.
    xhtml = zip_text(epub)
    assert "<strong>жирный</strong>" in xhtml and "<em>курсив</em>" in xhtml
    assert 'href="https://ru.wikipedia.org/"' in xhtml

    back = run(CONV, epub, tmp_path / "back.fb2", "epub", "fb2")
    fb = fb2_info(back)
    assert fb["title"] == "Война и мир"
    assert fb["author"] == "Лев Толстой"
    assert fb["lang"] == "ru"
    assert [s.split(" ")[:2] for s in fb["sections"]] == [
        ["Глава", "первая"], ["Глава", "вторая"], ["Глава", "третья"]]
    assert fb["cover"] and fb["cover"][1:] in fb["binaries"]
    assert len(fb["binaries"]) == 2
    assert fb["binaries"][fb["cover"][1:]][:2] == b"\xff\xd8"  # обложка — тот же JPEG
    assert any(data.startswith(b"\x89PNG") for data in fb["binaries"].values())
    raw = back.read_text("utf-8")
    assert "<strong>жирный</strong>" in raw and "<emphasis>курсив</emphasis>" in raw
    assert detect(back).fmt == "fb2"


@pytest.mark.parametrize("version", ["2.0", "3.0"])
def test_p18_epub_reading_spine_toc_cover(tmp_path, inproc, version):
    src = make_epub(tmp_path / "книга.epub", version=version)
    fb = fb2_info(run(CONV, src, tmp_path / "out.fb2", "epub", "fb2"))
    assert fb["title"] == "Тестовая книга"
    assert fb["author"] == "Фёдор Достоевский"
    # Порядок — по spine (manifest записан в обратном порядке), страница-обложка
    # отдельной главой не стала, подпись — из заголовка главы.
    assert [s.split(" ")[0] for s in fb["sections"]] == ["Начало", "Середина", "Конец"]
    assert "Вторая глава книги" in fb["sections"][1]
    assert fb["cover"] is not None
    assert len(fb["binaries"]) == 2
    epub = run(CONV, tmp_path / "out.fb2", tmp_path / "again.epub", "fb2", "epub")
    info = epub_info(epub)
    assert info["title"] == "Тестовая книга"
    _order(" ".join(info["chapters"]), "Первая глава книги", "Вторая глава", "Третья глава")


def test_epub_toc_labels_used_when_no_heading(tmp_path, inproc):
    chapters = [("Альфа", "<p>Текст без заголовка один.</p>"),
                ("Бета", "<p>Текст без заголовка два.</p>")]
    src = make_epub(tmp_path / "b.epub", version="3.0", chapters=chapters, cover=False)
    fb = fb2_info(run(CONV, src, tmp_path / "o.fb2", "epub", "fb2"))
    assert fb["sections"][0].startswith("Навигация Альфа")
    assert fb["sections"][1].startswith("Навигация Бета")


# --- П18: MOBI ---------------------------------------------------------------------------


@pytest.mark.parametrize("compression,flags", [(1, 0), (2, 0), (2, 3)])
def test_p18_mobi_to_epub_fb2_txt(tmp_path, inproc, compression, flags):
    src = make_mobi(tmp_path / "книга.mobi", compression=compression, extra_flags=flags,
                    html=MOBI_HTML.replace("</body>", "<p>" + "Длинный текст для нескольких записей. " * 300 + "</p></body>"))
    assert detect(src).fmt == "mobi"
    info = epub_info(run(CONV, src, tmp_path / "o.epub", "mobi", "epub"))
    assert info["title"] == "Мобильная книга"
    assert info["authors"] == ["Антон Чехов"]
    assert len(info["chapters"]) == 4  # обложка + 3 главы по mbp:pagebreak
    _order(" ".join(info["chapters"]), "Первая глава", "ёжик", "Вторая глава",
           "Третья глава", "Конец книги", "Длинный текст")
    assert len(info["cover"]) == 1 and len(info["images"]) == 2
    assert "alert" not in zip_text(tmp_path / "o.epub")
    fb = fb2_info(run(CONV, src, tmp_path / "o.fb2", "mobi", "fb2"))
    assert [s.split(" ")[:2] for s in fb["sections"]] == [
        ["Первая", "глава"], ["Вторая", "глава"], ["Третья", "глава"]]
    assert fb["cover"] is not None
    txt = run(CONV, src, tmp_path / "o.txt", "mobi", "txt").read_bytes()
    assert not txt.startswith(b"\xef\xbb\xbf") and b"\r\n" in txt
    text = txt.decode("utf-8")
    assert text.startswith("Мобильная книга\r\nАнтон Чехов")
    _order(text, "Первая глава", "Текст первой главы, ёжик.", "Вторая глава", "Третья глава")
    assert text.count("Длинный текст для нескольких записей.") == 300


def test_lz77_roundtrip():
    data = ("Привет, мир! " * 50 + "abc abc abc \x00\x01\x02 конец").encode("utf-8")
    assert books.palmdoc_decompress(lz77(data)) == data


@pytest.mark.parametrize("kwargs,message", [
    ({"compression": 17480}, "HUFF/CDIC"),
    ({"encryption": 2}, "DRM"),
    ({"encryption": 1}, "DRM"),
    ({"version": 8}, "KF8"),
    ({"compression": 7}, "неизвестный способ сжатия"),
])
def test_mobi_unsupported_variants(tmp_path, inproc, kwargs, message):
    src = make_mobi(tmp_path / "x.mobi", **kwargs)
    with pytest.raises(ConversionError, match="Этот вариант MOBI не поддерживается") as info:
        run(CONV, src, tmp_path / "o.epub", "mobi", "epub")
    assert message in str(info.value)


def test_mobi_broken_records(tmp_path, inproc):
    good = make_mobi(tmp_path / "g.mobi").read_bytes()
    bad = bytearray(good)
    struct.pack_into(">I", bad, 78 + 8, len(good) + 100)  # смещение за концом файла
    (tmp_path / "b.mobi").write_bytes(bytes(bad))
    with pytest.raises(ConversionError, match="повреждён"):
        run(CONV, tmp_path / "b.mobi", tmp_path / "o.txt", "mobi", "txt")
    # Обратная ссылка LZ77 дальше начала записи.
    with pytest.raises(_isolated.TaskError):
        books.palmdoc_decompress(b"\x80\x50")


# --- П18: книги → TXT / PDF ------------------------------------------------------------


def test_p18_books_to_txt(tmp_path, inproc):
    txt = run(CONV, make_fb2(tmp_path / "к.fb2"), tmp_path / "к.txt", "fb2", "txt")
    text = txt.read_bytes().decode("utf-8")
    assert text.startswith("Война и мир\r\nЛев Толстой")
    _order(text, "Глава первая", "Генуя", "Второй жирный и курсив абзац", "Глава вторая",
           "Анна Павловна", "Глава третья")
    text2 = run(CONV, make_epub(tmp_path / "e.epub"), tmp_path / "e.txt", "epub", "txt").read_text(
        "utf-8")
    _order(text2, "Тестовая книга", "Начало", "Первая глава книги", "Середина", "Конец")


def test_p18_books_to_pdf(tmp_path, inproc):
    out = run(CONV, make_fb2(tmp_path / "к.fb2"), tmp_path / "к.pdf", "fb2", "pdf")
    doc = pdfium.PdfDocument(str(out))
    pages = len(doc)
    doc.close()
    assert pages == 4  # обложка + 3 главы, каждая с новой страницы
    text = pdf_text(out)
    _order(text, "Глава первая", "Генуя и Лукка", "Глава вторая", "Анна Павловна", "Глава третья")
    out2 = run(CONV, make_mobi(tmp_path / "м.mobi"), tmp_path / "м.pdf", "mobi", "pdf",
               Preset.SMALL)
    _order(pdf_text(out2), "Первая глава", "ёжик", "Третья глава")


def test_pdf_long_words_and_big_text(tmp_path, inproc):
    long_word = "Ж" * 3000
    paras = [long_word, "обычный " * 2000] + [f"Абзац номер {i}" for i in range(400)]
    src = make_fb2(tmp_path / "длинная.fb2", chapters=[("Одна", paras)], cover=False, image=False)
    out = run(CONV, src, tmp_path / "д.pdf", "fb2", "pdf")
    text = pdf_text(out)
    assert "Абзац номер 399" in text
    assert text.count("Ж") >= 2990


def test_pdf_without_fonts_is_error(tmp_path, inproc, monkeypatch):
    empty = tmp_path / "нет шрифтов"
    empty.mkdir()
    monkeypatch.setitem(_docjob.LIMITS, "fonts_dir", str(empty))
    with pytest.raises(ConversionError, match="Не найден шрифт"):
        run(CONV, make_fb2(tmp_path / "к.fb2"), tmp_path / "к.pdf", "fb2", "pdf")


# --- Безопасность: XXE, «миллиард смешков» ------------------------------------------


def test_xxe_in_fb2(tmp_path, inproc, secret):
    doctype = f'<!DOCTYPE FictionBook [<!ENTITY xxe SYSTEM "{file_uri(secret)}">]>'
    src = make_fb2(tmp_path / "xxe.fb2", doctype=doctype,
                   chapters=[("Глава", ["До &xxe; после"])], cover=False, image=False)
    for tid in ("epub", "txt"):
        out = tmp_path / f"o.{tid}"
        with pytest.raises(ConversionError, match="сущност"):
            run(CONV, src, out, "fb2", tid)
        assert b"\xd0\xa1\xd0\x95\xd0\x9a" not in out.read_bytes()  # «СЕК»


def test_billion_laughs_fb2_fast(tmp_path, inproc):
    src = make_fb2(tmp_path / "lol.fb2", doctype=LAUGHS, chapters=[("Г", ["&lol9;"])],
                   cover=False, image=False)
    t0 = time.monotonic()
    with pytest.raises(ConversionError, match="сущност"):
        run(CONV, src, tmp_path / "o.txt", "fb2", "txt")
    assert time.monotonic() - t0 < 5


def test_xxe_in_epub_chapter_and_opf(tmp_path, inproc, secret):
    doctype = f'<!DOCTYPE html [<!ENTITY xxe SYSTEM "{file_uri(secret)}">]>'
    chapters = [("Г", "<p>До &xxe; после</p>")]
    src = make_epub(tmp_path / "x1.epub", chapters=chapters, chapter_doctype=doctype)
    out = tmp_path / "o1.txt"
    with pytest.raises(ConversionError, match="сущност"):
        run(CONV, src, out, "epub", "txt")
    assert "СЕКРЕТ" not in out.read_text("utf-8", errors="ignore")
    src2 = make_epub(tmp_path / "x2.epub", opf_doctype=LAUGHS.replace("lolz", "package"))
    with pytest.raises(ConversionError, match="сущност"):
        run(CONV, src2, tmp_path / "o2.txt", "epub", "txt")
    # Parameter entity (внешний DTD через %…;) тоже не читается.
    pe = f'<!DOCTYPE html [<!ENTITY % ext SYSTEM "{file_uri(secret)}"> %ext;]>'
    src3 = make_epub(tmp_path / "x3.epub", chapters=[("Г", "<p>Текст</p>")], chapter_doctype=pe)
    out3 = tmp_path / "o3.txt"
    try:
        run(CONV, src3, out3, "epub", "txt")
    except ConversionError:
        pass
    assert "СЕКРЕТ" not in out3.read_text("utf-8", errors="ignore")


def test_xhtml_with_html_entities_falls_back_to_html_parser(tmp_path, inproc):
    doctype = ('<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.1//EN" '
               '"http://www.w3.org/TR/xhtml11/DTD/xhtml11.dtd">')
    chapters = [("Г", "<p>Неразрывный&nbsp;пробел &mdash; тире</p>")]
    src = make_epub(tmp_path / "e.epub", chapters=chapters, chapter_doctype=doctype, cover=False)
    text = run(CONV, src, tmp_path / "o.txt", "epub", "txt").read_text("utf-8")
    assert "Неразрывный\xa0пробел — тире" in text


# --- Безопасность: ZIP-бомбы ---------------------------------------------------------------


def _epub_members(tmp_path: Path) -> dict[str, bytes]:
    src = make_epub(tmp_path / "base.epub")
    with zipfile.ZipFile(src) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_zip_bomb_entry_in_epub(tmp_path, inproc):
    members = _epub_members(tmp_path)
    members["OPS/text/c0.xhtml"] = b""  # заменим «бомбой» ниже
    del members["OPS/text/c0.xhtml"]
    src = bomb_zip(tmp_path / "bomb.epub", members=members, bomb_name="OPS/text/c0.xhtml")
    t0 = time.monotonic()
    with pytest.raises(ConversionError, match="ZIP-бомб|Слишком большой объём"):
        run(CONV, src, tmp_path / "o.txt", "epub", "txt")
    assert time.monotonic() - t0 < 10


def test_zip_false_size_header_in_epub(tmp_path, inproc):
    members = _epub_members(tmp_path)
    members["OPS/text/c1.xhtml"] = (b"<html xmlns='http://www.w3.org/1999/xhtml'><body><p>"
                                    + b"a" * (40 * 2**20) + b"</p></body></html>")
    src = tmp_path / "lie.epub"
    with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data, compress_type=zipfile.ZIP_STORED if name == "mimetype"
                             else zipfile.ZIP_DEFLATED)
    lie_about_size(src, "OPS/text/c1.xhtml", 100)
    t0 = time.monotonic()
    with pytest.raises(ConversionError):
        run(CONV, src, tmp_path / "o.txt", "epub", "txt")
    assert time.monotonic() - t0 < 10


def test_zip_many_entries_epub(tmp_path, inproc):
    src = many_entries_zip(tmp_path / "many.epub", first=_epub_members(tmp_path))
    t0 = time.monotonic()
    with pytest.raises(ConversionError, match="Слишком много частей"):
        run(CONV, src, tmp_path / "o.txt", "epub", "txt")
    assert time.monotonic() - t0 < 10


def test_zip_total_limit(tmp_path, inproc, monkeypatch):
    monkeypatch.setitem(_docjob.LIMITS, "zip_max_total", 2000)
    with pytest.raises(ConversionError, match="Слишком большой объём"):
        run(CONV, make_epub(tmp_path / "e.epub"), tmp_path / "o.txt", "epub", "txt")


def _noise_png(side: int = 400) -> bytes:
    """PNG из случайных пикселей (почти не сжимается): side=400 → ≈480 КБ."""
    import io
    import random

    from PIL import Image

    rng = random.Random(1)
    image = Image.frombytes("RGB", (side, side), bytes(rng.randrange(256)
                                                        for _ in range(side * side * 3)))
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def test_epub_repeated_image_read_once(tmp_path, inproc, monkeypatch):
    """Одна картинка во многих главах распаковывается один раз (нет ложного zip_too_big)."""
    big = _noise_png()
    assert len(big) > 400_000
    chapters = [(f"Г{i}", f'<p>Глава {i}.</p><p><img src="../images/big.png"/></p>')
                for i in range(30)]
    src = make_epub(tmp_path / "e.epub", chapters=chapters, cover=False,
                    extra={"OPS/images/big.png": big})
    # Всего в архиве ≈0,5 МБ текста и картинок; 30 повторных распаковок — ≈14 МБ.
    monkeypatch.setitem(_docjob.LIMITS, "zip_max_total", 4 * 2**20)
    fb2 = run(CONV, src, tmp_path / "o.fb2", "epub", "fb2").read_text("utf-8")
    assert "Глава 29." in fb2
    assert fb2.count("<binary ") == 1


def test_epub_drm_and_paths_outside(tmp_path, inproc, secret):
    enc = ('<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
           '<EncryptedData xmlns="http://www.w3.org/2001/04/xmlenc#"><CipherData>'
           '<CipherReference URI="OPS/text/c0.xhtml"/></CipherData></EncryptedData></encryption>')
    src = make_epub(tmp_path / "drm.epub", extra={"META-INF/encryption.xml": enc})
    with pytest.raises(ConversionError, match="DRM"):
        run(CONV, src, tmp_path / "o.txt", "epub", "txt")
    # Ссылки на картинки вне архива и абсолютные — не читаются (секрет не попадает).
    shutil_copy = tmp_path / "pic.png"
    shutil_copy.write_bytes(png_bytes())
    body = (f'<p>Т</p><img src="../../../../{secret.name}"/><img src="{file_uri(shutil_copy)}"/>'
            f'<img src="/etc/passwd"/><img src="C:\\Windows\\win.ini"/>')
    src2 = make_epub(tmp_path / "out.epub", chapters=[("Г", body)], cover=False)
    fb = fb2_info(run(CONV, src2, tmp_path / "o.fb2", "epub", "fb2"))
    assert fb["binaries"] == {}


# --- Безопасность: активное содержимое в EPUB-выходе ----------------------------------


def test_sanitizer_in_epub_output(tmp_path, inproc):
    # Невалидный XML (незакрытые <img>, <input>) — глава разбирается HTML-парсером.
    src = make_epub(tmp_path / "evil.epub", chapters=[("Злая", EVIL_HTML)], cover=False)
    out = run(CONV, src, tmp_path / "o.fb2", "epub", "fb2")
    raw = out.read_text("utf-8")
    assert_clean(raw)
    assert "https://example.com/путь" in raw
    epub = run(CONV, out, tmp_path / "again.epub", "fb2", "epub")
    text = zip_text(epub)
    assert_clean(text.replace('href="../style.css"', ""))
    assert "<strong>жирный</strong>" in text
    assert 'href="https://example.com/путь"' in text
    assert "Пункт один" in text and "Иван" in text


def test_sanitizer_mobi_to_epub(tmp_path, inproc):
    src = make_mobi(tmp_path / "evil.mobi", "<html><body>" + EVIL_HTML + "</body></html>")
    out = run(CONV, src, tmp_path / "o.epub", "mobi", "epub")
    text = zip_text(out)
    assert_clean(text)
    assert "<ul>" in text and "<table>" in text and "<h1>" in text
    assert 'href="https://example.com/путь"' in text


# --- Битые, пустые, мусорные файлы ------------------------------------------------------


@pytest.mark.parametrize("fmt", ["epub", "fb2", "mobi"])
@pytest.mark.parametrize("payload", [b"\x00" * 10, b"PK\x03\x04garbage" * 50,
                                     b"<FictionBook></FictionBook>", os.urandom(4096)])
def test_garbage_is_conversion_error(tmp_path, inproc, fmt, payload):
    src = tmp_path / f"x.{fmt}"
    src.write_bytes(payload)
    with pytest.raises(ConversionError):
        run(CONV, src, tmp_path / "o.txt", fmt, "txt")


def test_empty_file_and_book_without_text(tmp_path, inproc):
    empty = tmp_path / "e.fb2"
    empty.write_bytes(b"")
    with pytest.raises(ConversionError, match="нет содержимого"):
        run(CONV, empty, tmp_path / "o.txt", "fb2", "txt")
    src = tmp_path / "n.fb2"
    src.write_text('<?xml version="1.0"?><FictionBook xmlns="http://www.gribuser.ru/xml/'
                   'fictionbook/2.0"><description/><body/></FictionBook>', encoding="utf-8")
    with pytest.raises(ConversionError, match="нет содержимого"):
        run(CONV, src, tmp_path / "o2.txt", "fb2", "txt")


def test_wrong_format_and_target_checks(tmp_path, inproc):
    src = make_fb2(tmp_path / "к.fb2")
    t = target(CONV, "fb2", "epub")
    (tmp_path / "o.epub").touch()
    with pytest.raises(ConversionError):
        CONV.convert([src], tmp_path / "o.epub", t, Preset.MEDIUM, lambda _p: None,
                     CancelToken(), formats=["fb2", "fb2"])
    with pytest.raises(ConversionError):
        CONV.convert([src], tmp_path / "o.epub", t, Preset.MEDIUM, lambda _p: None,
                     CancelToken(), formats=["epub"])  # «в тот же формат» не предлагается


def test_too_many_chapters_limit(tmp_path, inproc, monkeypatch):
    monkeypatch.setitem(_docjob.LIMITS, "max_chapters", 2)
    with pytest.raises(ConversionError, match="Слишком много глав"):
        run(CONV, make_fb2(tmp_path / "к.fb2"), tmp_path / "o.txt", "fb2", "txt")


# --- Настоящий дочерний процесс: работа, отмена, JobManager -----------------------------


def test_real_child_fb2_to_epub_and_unicode_long_path(tmp_path):
    base = naming.os_path(tmp_path / ("д" * 120) / ("е" * 120))
    os.makedirs(base)
    src = Path(os.path.join(base, "книга с пробелами — №1.fb2"))
    make_fb2(src)
    out = Path(os.path.join(base, ".~к.part.epub"))
    open(out, "wb").close()
    assert len(str(out)) > 260
    CONV.convert([src], out, target(CONV, "fb2", "epub"), Preset.MEDIUM, lambda _p: None,
                 CancelToken(), formats=["fb2"])
    assert epub_info(out)["title"] == "Война и мир"
    assert no_children()


def test_cancel_long_book_to_pdf(tmp_path):
    paras = [f"Абзац {i}: " + "слово " * 60 for i in range(6000)]
    src = make_fb2(tmp_path / "толстая.fb2", chapters=[(f"Глава {c}", paras[c::10])
                                                       for c in range(10)],
                   cover=False, image=False)
    token = CancelToken()
    errors: list[BaseException] = []
    progressed = threading.Event()

    def work():
        try:
            run(CONV, src, tmp_path / "o.pdf", "fb2", "pdf",
                progress=lambda _p: progressed.set(), cancel=token)
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)

    thread = threading.Thread(target=work)
    thread.start()
    assert progressed.wait(60)
    t0 = time.monotonic()
    token.cancel()
    thread.join(10)
    assert time.monotonic() - t0 < 2.0
    assert errors and isinstance(errors[0], Cancelled)
    assert no_children()


def test_job_manager_end_to_end(tmp_path):
    folder = tmp_path / "книги"
    folder.mkdir()
    src = make_fb2(folder / "моя книга.fb2")
    manager = JobManager(max_workers=1)
    try:
        job = manager.submit([src], CONV, target(CONV, "fb2", "epub"), Preset.MEDIUM, folder,
                             formats=["fb2"])
        wait_jobs([job])
    finally:
        manager.shutdown(timeout=10)
    assert job.status == JobStatus.DONE, job.error_message
    assert job.result_path == folder / "моя книга.epub"
    assert epub_info(job.result_path)["authors"] == ["Лев Толстой"]
    assert sorted(p.name for p in folder.iterdir()) == ["моя книга.epub", "моя книга.fb2"]
    assert no_children()


def test_docs_task_and_pool(monkeypatch):
    assert "docs" in _isolated.POOLS
    assert _isolated.TASKS["docs"].startswith("konverter.converters._docjob:")
    monkeypatch.setattr(_isolated, "DOCS_MAX_CHILDREN", 1)
    assert _isolated._slot_limit("docs") == 1


# --- Метаданные картинок (EXIF/GPS, тексты PNG) ----------------------------------------------

SECRET_MARK = b"SECRET-"


def _segment(marker: int, payload: bytes) -> bytes:
    return bytes([0xFF, marker]) + struct.pack(">H", len(payload) + 2) + payload


def dirty_jpeg() -> tuple[bytes, bytes]:
    """JPEG с EXIF (GPS, описание), XMP, IPTC (APP13), COM и ICC; (байты, профиль ICC)."""
    import io

    from PIL import Image, ImageCms

    icc = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    image = Image.new("RGB", (64, 48))
    image.putdata([((x * 4) % 256, (y * 5) % 256, (x * y) % 256)
                   for y in range(48) for x in range(64)])
    exif = Image.Exif()
    exif[0x010E] = "SECRET-DESCRIPTION"
    exif[0x013B] = "SECRET-ARTIST"
    gps = exif.get_ifd(0x8825)
    gps[1] = "N"
    gps[2] = (55.0, 45.0, 21.0)
    gps[3] = "E"
    gps[4] = (37.0, 37.0, 4.0)
    out = io.BytesIO()
    image.save(out, format="JPEG", quality=90, exif=exif, icc_profile=icc)
    data = out.getvalue()
    extra = (_segment(0xE1, b"http://ns.adobe.com/xap/1.0/\x00<x:xmpmeta>SECRET-XMP</x:xmpmeta>")
             + _segment(0xED, b"Photoshop 3.0\x008BIM\x04\x04\x00\x00SECRET-IPTC")
             + _segment(0xFE, b"SECRET-COMMENT")
             + _segment(0xEE, b"Adobe\x00d\x00\x00\x00\x00\x01"))
    assert b"GPS" in data or exif.get_ifd(0x8825)
    return data[:2] + extra + data[2:], icc


def _png_chunk(kind: bytes, payload: bytes) -> bytes:
    import zlib

    return (struct.pack(">I", len(payload)) + kind + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF))


def dirty_png() -> bytes:
    """PNG с tEXt, zTXt, iTXt, eXIf и tIME."""
    import io

    from PIL import Image, PngImagePlugin

    image = Image.new("RGB", (40, 30))
    image.putdata([((x * 6) % 256, (y * 8) % 256, 90) for y in range(30) for x in range(40)])
    info = PngImagePlugin.PngInfo()
    info.add_text("Comment", "SECRET-TEXT")
    info.add_text("Author", "SECRET-ZTXT", zip=True)
    info.add_itxt("Description", "SECRET-ITXT", lang="ru", tkey="Описание")
    exif = Image.Exif()
    exif[0x013B] = "SECRET-PNG-EXIF"
    out = io.BytesIO()
    image.save(out, format="PNG", pnginfo=info, exif=exif)
    data = out.getvalue()
    iend = data.rindex(b"IEND") - 4
    data = data[:iend] + _png_chunk(b"tIME", b"\x07\xe8\x01\x02\x03\x04\x05") + data[iend:]
    for kind in (b"tEXt", b"zTXt", b"iTXt", b"eXIf", b"tIME"):
        assert kind in data
    return data


def _pixels(data: bytes):
    import io

    from PIL import Image

    with Image.open(io.BytesIO(data)) as image:
        return image.convert("RGB").tobytes(), image.size


def _assert_clean_image(data: bytes, original: bytes) -> None:
    import io

    from PIL import Image

    assert SECRET_MARK not in data
    for bad in (b"Exif\x00\x00", b"eXIf", b"tEXt", b"zTXt", b"iTXt", b"tIME", b"Photoshop 3.0",
                b"ns.adobe.com/xap"):
        assert bad not in data, bad
    with Image.open(io.BytesIO(data)) as image:
        assert not image.getexif()
        assert "comment" not in image.info and not getattr(image, "text", {})
    assert _pixels(data) == _pixels(original)


def _images_from_output(path: Path) -> list[bytes]:
    import re

    data = path.read_bytes()
    suffix = path.suffix.lower()
    if suffix in (".epub", ".docx"):
        with zipfile.ZipFile(path) as archive:
            return [archive.read(n) for n in archive.namelist()
                    if n.lower().endswith((".jpg", ".jpeg", ".png"))
                    and "thumbnail" not in n]
    text = data.decode("utf-8")
    if suffix == ".fb2":
        found = re.findall(r"<binary[^>]*>([^<]*)</binary>", text)
    else:
        found = re.findall(r'src="data:image/[a-z]+;base64,([^"]+)"', text)
    return [base64.b64decode(item) for item in found]


def _check_output(path: Path, originals: list[bytes]) -> None:
    images = _images_from_output(path)
    assert len(images) == len(originals), (path.name, len(images))
    assert SECRET_MARK not in path.read_bytes()
    by_size = {_pixels(o)[1]: o for o in originals}
    for image in images:
        _assert_clean_image(image, by_size[_pixels(image)[1]])


def test_image_metadata_stripped_everywhere(tmp_path, inproc):
    from konverter.converters import documents

    jpeg, icc = dirty_jpeg()
    png = dirty_png()
    originals = [jpeg, png]
    # EPUB → FB2
    chapters = [("Глава", '<p>Текст.</p><p><img src="../images/a.jpg"/></p>'
                          '<p><img src="../images/b.png"/></p>')]
    src = make_epub(tmp_path / "e.epub", chapters=chapters, cover=False,
                    extra={"OPS/images/a.jpg": jpeg, "OPS/images/b.png": png})
    _check_output(run(CONV, src, tmp_path / "o.fb2", "epub", "fb2"), originals)
    # FB2 → EPUB
    fb2 = make_fb2(tmp_path / "k.fb2")
    text = fb2.read_text("utf-8")
    text = text.replace(base64.b64encode(jpeg_bytes()).decode(), base64.b64encode(jpeg).decode())
    text = text.replace(base64.b64encode(png_bytes()).decode(), base64.b64encode(png).decode())
    fb2.write_text(text, encoding="utf-8")
    epub = run(CONV, fb2, tmp_path / "o.epub", "fb2", "epub")
    _check_output(epub, originals)
    # ICC и APP14 (Adobe) у JPEG сохраняются.
    with zipfile.ZipFile(epub) as archive:
        kept = next(archive.read(n) for n in archive.namelist() if n.endswith(".jpg"))
    assert icc[:64] in kept and b"Adobe\x00" in kept
    # HTML (data:) → DOCX и MD (сырой <img>) → HTML.
    tags = "".join(f'<p><img src="data:image/{kind};base64,{base64.b64encode(raw).decode()}"></p>'
                   for kind, raw in (("jpeg", jpeg), ("png", png)))
    html_src = tmp_path / "d.html"
    html_src.write_text(f"<html><body><p>Текст</p>{tags}</body></html>", encoding="utf-8")
    docs = documents.DocumentConverter()
    _check_output(run(docs, html_src, tmp_path / "o.docx", "html", "docx"), originals)
    md_src = tmp_path / "d.md"
    md_src.write_text("# Заголовок\n\n" + tags + "\n", encoding="utf-8")
    _check_output(run(docs, md_src, tmp_path / "o.html", "md", "html"), originals)


@pytest.mark.parametrize("payload", [
    b"\xff\xd8\xff\xe1\xff\xff" + b"x" * 10,  # длина сегмента за концом файла
    b"\xff\xd8\xff\xe1\x00\x01",  # длина меньше 2
    b"\xff\xd8" + b"\x00" * 100,  # нет маркера
    b"\x89PNG\r\n\x1a\n" + b"\xff\xff\xff\xff" + b"IHDR",  # длина чанка за концом
    b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\x0dIHDR",  # обрыв
])
def test_strip_metadata_damaged_structure(payload):
    from konverter.converters._docmodel import strip_jpeg_metadata, strip_png_metadata

    strip = strip_jpeg_metadata if payload.startswith(b"\xff\xd8") else strip_png_metadata
    assert strip(payload) is None


def test_strip_metadata_fallback_reencode(monkeypatch):
    """Разбор не удался — картинка перекодируется Pillow без метаданных."""
    from konverter.converters import _docmodel

    jpeg, _icc = dirty_jpeg()
    monkeypatch.setattr(_docmodel, "strip_jpeg_metadata", lambda data: None)
    picture = _docmodel.load_picture(jpeg, 10**7)
    assert picture is not None and SECRET_MARK not in picture.data
    assert b"Exif\x00\x00" not in picture.data
    _pixels(picture.data)  # декодируется


def test_strip_jpeg_camera_and_progressive():
    """Камерный JPEG (EXIF вместо JFIF) и прогрессивный: пиксели те же, python-docx принимает."""
    import io

    import docx
    from PIL import Image

    from konverter.converters._docmodel import strip_jpeg_metadata

    jpeg, _icc = dirty_jpeg()
    with Image.open(io.BytesIO(jpeg)) as image:
        out = io.BytesIO()
        image.save(out, format="JPEG", quality=85, progressive=True, exif=image.getexif())
        progressive = out.getvalue()
    for source in (jpeg, progressive):
        data = bytearray(source)
        # Убрать APP0 JFIF, если есть: как у камеры — сразу APP1 Exif.
        if data[2:4] == b"\xff\xe0":
            length = struct.unpack(">H", data[4:6])[0]
            del data[2:4 + length]
        data = bytes(data)
        clean = strip_jpeg_metadata(data)
        assert clean is not None and clean[6:10] == b"JFIF"
        assert b"Exif\x00\x00" not in clean and SECRET_MARK not in clean
        assert _pixels(clean) == _pixels(data)
        document = docx.Document()
        document.add_picture(io.BytesIO(clean))
    assert strip_jpeg_metadata(progressive + b"TRAILER-SECRET-") == strip_jpeg_metadata(progressive)


def test_load_picture_applies_exif_orientation():
    """Orientation=6 (камера повёрнута): пиксели поворачиваются, EXIF убирается."""
    import io

    from PIL import Image

    from konverter.converters._docmodel import load_picture

    image = Image.new("RGB", (60, 20), (200, 10, 10))
    exif = Image.Exif()
    exif[0x0112] = 6
    gps = exif.get_ifd(0x8825)
    gps[1] = "N"
    gps[2] = (55.0, 45.0, 21.0)
    out = io.BytesIO()
    image.save(out, format="JPEG", exif=exif)
    picture = load_picture(out.getvalue(), 10**7)
    assert picture is not None and (picture.width, picture.height) == (20, 60)
    assert b"Exif\x00\x00" not in picture.data
    with Image.open(io.BytesIO(picture.data)) as result:
        assert result.size == (20, 60) and not result.getexif()


# --- Белые списки чанков PNG и маркеров JPEG, MPO, битые JPEG ---------------------------------


def _test_jpeg(**save) -> bytes:
    import io

    from PIL import Image

    image = Image.new("RGB", (64, 48))
    image.putdata([((x * 4) % 256, (y * 5) % 256, (x * y) % 256)
                   for y in range(48) for x in range(64)])
    out = io.BytesIO()
    image.save(out, format="JPEG", quality=90, **save)
    return out.getvalue()


def _png_chunk(kind: bytes, payload: bytes) -> bytes:
    import zlib

    return (struct.pack(">I", len(payload)) + kind + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF))


def _png_chunks(data: bytes) -> list[bytes]:
    kinds, pos = [], 8
    while pos < len(data):
        length = struct.unpack(">I", data[pos:pos + 4])[0]
        kinds.append(data[pos + 4:pos + 8])
        pos += 12 + length
    return kinds


def _jpeg_markers(data: bytes) -> list[int]:
    """Маркеры сегментов до первого SOS (с длиной)."""
    markers, pos = [], 2
    while data[pos] == 0xFF and data[pos + 1] != 0xDA:
        markers.append(data[pos + 1])
        pos += 2 + struct.unpack(">H", data[pos + 2:pos + 4])[0]
    return markers


def test_strip_png_whitelist():
    """caBX (C2PA), zxIf, выдуманный prVt, tEXt вырезаются; пиксели те же."""
    import io

    from PIL import Image

    from konverter.converters._docmodel import load_picture, strip_png_metadata

    image = Image.new("RGBA", (40, 30))
    image.putdata([(x * 6, y * 8, (x + y) % 256, 200) for y in range(30) for x in range(40)])
    out = io.BytesIO()
    image.save(out, format="PNG")
    original = out.getvalue()
    iend = original.rfind(b"IEND") - 4
    extra = b"".join(_png_chunk(kind, SECRET_MARK + kind) for kind in
                     (b"caBX", b"zxIf", b"prVt", b"tEXt")) + _png_chunk(b"pHYs", b"\0\0\x0b\x13" * 2 + b"\1")
    ihdr_end = 8 + 12 + 13
    dirty = original[:ihdr_end] + extra + original[ihdr_end:iend] + extra + original[iend:]
    allowed = {b"IHDR", b"PLTE", b"IDAT", b"IEND", b"tRNS", b"gAMA", b"cHRM", b"sRGB",
               b"iCCP", b"sBIT", b"pHYs", b"bKGD", b"acTL", b"fcTL", b"fdAT"}
    clean = strip_png_metadata(dirty)
    assert clean is not None and SECRET_MARK not in clean
    assert set(_png_chunks(clean)) <= allowed and b"pHYs" in _png_chunks(clean)
    assert _pixels(clean) == _pixels(original)
    picture = load_picture(dirty, 10**7)
    assert picture is not None and picture.data == clean


def test_strip_png_keeps_apng_frames():
    import io

    from PIL import Image

    from konverter.converters._docmodel import strip_png_metadata

    frames = [Image.new("RGB", (8, 8), color) for color in ((255, 0, 0), (0, 255, 0))]
    out = io.BytesIO()
    frames[0].save(out, format="PNG", save_all=True, append_images=frames[1:])
    clean = strip_png_metadata(out.getvalue())
    assert clean is not None
    assert {b"acTL", b"fcTL", b"fdAT"} <= set(_png_chunks(clean))


def test_load_picture_mpo_as_jpeg():
    """JPEG с MPF (Pillow: «MPO») — как JPEG: без перекодирования, одна картинка, без MPF."""
    import io

    from PIL import Image

    from konverter.converters._docmodel import load_picture

    first = Image.open(io.BytesIO(_test_jpeg())).convert("RGB")
    second = Image.new("RGB", (64, 48), (10, 200, 30))
    out = io.BytesIO()
    first.save(out, format="MPO", save_all=True, append_images=[second], quality=90)
    mpo = out.getvalue()
    with Image.open(io.BytesIO(mpo)) as probe:
        assert probe.format == "MPO" and probe.n_frames == 2
    picture = load_picture(mpo, 10**7)
    assert picture is not None and picture.media_type == "image/jpeg"
    assert len(picture.data) <= len(mpo)
    assert b"MPF\x00" not in picture.data and picture.data.count(b"\xff\xd8") == 1
    with Image.open(io.BytesIO(picture.data)) as result:
        assert result.format == "JPEG" and getattr(result, "n_frames", 1) == 1
    with Image.open(io.BytesIO(mpo)) as source:
        source.seek(0)
        expected = source.convert("RGB").tobytes()
    assert _pixels(picture.data)[0] == expected


def test_strip_jpeg_app0_jfif_trimmed_jfxx_dropped():
    from konverter.converters._docmodel import strip_jpeg_metadata

    original = _test_jpeg()
    assert original[2:4] == b"\xff\xe0"
    app0_end = 4 + struct.unpack(">H", original[4:6])[0]
    # JFIF 1.02, 72×72 dpi, миниатюра 2×2 (12 байт RGB) + хвост с «секретом».
    jfif = _segment(0xE0, b"JFIF\x00\x01\x02\x01\x00\x48\x00\x48\x02\x02"
                    + b"\x7f" * 12 + SECRET_MARK + b"THUMB")
    jfxx = _segment(0xE0, b"JFXX\x00\x10" + b"\xff\xd8" + SECRET_MARK + b"JFXX\xff\xd9")
    dirty = original[:2] + jfif + jfxx + original[app0_end:]
    clean = strip_jpeg_metadata(dirty)
    assert clean is not None and SECRET_MARK not in clean and b"JFXX" not in clean
    assert clean[2:20] == b"\xff\xe0\x00\x10JFIF\x00\x01\x02\x01\x00\x48\x00\x48\x00\x00"
    assert _jpeg_markers(clean).count(0xE0) == 1
    assert _pixels(clean) == _pixels(original)


def test_strip_jpeg_unknown_markers_dropped():
    """Зарезервированные и JPGn-маркеры (0x02–0xBF, 0xC8, 0xF0–0xFD) вырезаются."""
    from konverter.converters._docmodel import strip_jpeg_metadata

    original = _test_jpeg()
    app0_end = 4 + struct.unpack(">H", original[4:6])[0]
    odd = (0x02, 0x50, 0xBF, 0xC8, 0xF0, 0xF7, 0xFD)
    junk = b"".join(_segment(marker, SECRET_MARK + bytes([marker])) for marker in odd)
    dirty = original[:app0_end] + junk + original[app0_end:]
    clean = strip_jpeg_metadata(dirty)
    assert clean is not None and SECRET_MARK not in clean
    assert not set(_jpeg_markers(clean)) & set(odd)
    assert set(_jpeg_markers(original)) == set(_jpeg_markers(clean))  # нужное — на месте
    assert _pixels(clean) == _pixels(original)


def test_load_picture_rejects_broken_jpeg():
    """Структура цела, но не декодируется (нет скана / испорчен SOS) — картинка отбрасывается."""
    from konverter.converters._docmodel import load_picture

    original = _test_jpeg()
    sos = original.find(b"\xff\xda")
    no_scan = original[:sos] + b"\xff\xd9"
    bad_sos = original[:sos + 4] + b"\x09" + original[sos + 5:]
    assert load_picture(original, 10**7) is not None
    for broken in (no_scan, bad_sos):
        assert load_picture(broken, 10**7) is None
