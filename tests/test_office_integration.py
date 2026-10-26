"""Интеграционные тесты с настоящим LibreOffice (этап 5, П22).

Пропускаются, если модуль не установлен в `data/modules/libreoffice`
(устанавливается самой программой: скачивание с сервера TDF, проверка
подписей, распаковка). Тестовые файлы генерируются: DOCX — python-docx,
XLSX — openpyxl, ODP — вручную (минимальный ODF); DOC/ODT/RTF/XLS/ODS/PPTX/PPT
получаются из них самим LibreOffice в фикстуре.

После каждой конвертации и после отмены проверяется, что в системе нет
процессов soffice из папки модуля (psutil, по пути exe).
"""

from __future__ import annotations

import threading
import time
import zipfile
from pathlib import Path

import psutil
import pytest

from konverter.converters import _office
from konverter.converters.base import CancelToken, Cancelled, Category, ConversionError, Preset
from konverter.converters.documents import DocumentConverter
from konverter.converters.presentations import PresentationConverter
from konverter.converters.tables import SpreadsheetConverter
from konverter.core import naming
from konverter.libreoffice import jobobject, module, soffice

pytestmark = [
    pytest.mark.skipif(not module.is_installed(), reason="модуль LibreOffice не установлен"),
    pytest.mark.libreoffice,
]

SECRET_AUTHOR = "Секретный Автор Тестов"


def soffice_processes() -> list[psutil.Process]:
    """Процессы soffice.exe/soffice.bin из папки модуля."""
    root = str(module.module_path().resolve()).casefold()
    found = []
    for proc in psutil.process_iter(["name", "exe"]):
        try:
            name = (proc.info.get("name") or "").casefold()
            exe = (proc.info.get("exe") or "").casefold()
        except (psutil.Error, OSError):
            continue
        if name.startswith("soffice") and exe.startswith(root):
            found.append(proc)
    return found


def wait_no_soffice(timeout: float = 10.0) -> list[psutil.Process]:
    deadline = time.monotonic() + timeout
    while True:
        left = soffice_processes()
        if not left or time.monotonic() > deadline:
            return left
        time.sleep(0.2)


def run_dirs() -> list[Path]:
    temp = module.temp_root()
    return [p for p in temp.iterdir() if p.name.startswith(soffice.RUN_PREFIX)]


@pytest.fixture(autouse=True)
def no_leftovers():
    assert wait_no_soffice(5) == []
    before = set(run_dirs())
    yield
    assert wait_no_soffice(10) == [], "остались процессы soffice"
    assert set(run_dirs()) - before == set(), "осталась папка работы LibreOffice"


# --- Генерация тестовых файлов -------------------------------------------------------


_ODP_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.2">
<office:body><office:presentation>
<draw:page draw:name="s1"><draw:frame svg:x="2cm" svg:y="2cm" svg:width="20cm" svg:height="3cm">
<draw:text-box><text:p>Первый слайд</text:p></draw:text-box></draw:frame></draw:page>
<draw:page draw:name="s2"><draw:frame svg:x="2cm" svg:y="2cm" svg:width="20cm" svg:height="3cm">
<draw:text-box><text:p>Второй слайд</text:p></draw:text-box></draw:frame></draw:page>
</office:presentation></office:body></office:document-content>
"""
_ODP_MANIFEST = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">
<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.presentation"/>
<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>
"""


def make_odp(path: Path) -> Path:
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.presentation")
        zf.writestr("META-INF/manifest.xml", _ODP_MANIFEST, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("content.xml", _ODP_CONTENT, compress_type=zipfile.ZIP_DEFLATED)
    return path


def make_docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.core_properties.author = SECRET_AUTHOR
    document.core_properties.title = "Заголовок теста"
    document.add_heading("Привет, LibreOffice", 1)
    document.add_paragraph("Текст документа — ёлка.")
    document.save(path)
    return path


def make_xlsx(path: Path) -> Path:
    import openpyxl

    book = openpyxl.Workbook()
    book.properties.creator = SECRET_AUTHOR
    sheet = book.active
    sheet.title = "Лист1"
    sheet.append(["Имя", "Число"])
    sheet.append(["Ёжик", 42])
    second = book.create_sheet("Второй")
    second.append(["второй лист"])
    book.save(path)
    return path


def lo_convert(source: Path, convert_to: str, ext: str, out_dir: Path) -> Path:
    """Фикстура: исходник в другой формат самим LibreOffice (минуя код программы)."""
    with soffice.work_dir(module.temp_root()) as work:
        profile = work / "profile"
        soffice.write_profile(profile)
        args = soffice.build_args(module.soffice_path(), profile, convert_to, work, source)
        result = jobobject.run(args, timeout=300, cwd=work)
        produced = work / f"{source.stem}.{ext}"
        assert result.returncode == 0 and produced.is_file(), convert_to
        target = out_dir / f"sample.{ext}"
        target.write_bytes(produced.read_bytes())
    return target


@pytest.fixture(scope="module")
def samples(tmp_path_factory) -> dict[str, Path]:
    folder = tmp_path_factory.mktemp("office_samples")
    files = {
        "docx": make_docx(folder / "sample.docx"),
        "xlsx": make_xlsx(folder / "sample.xlsx"),
        "odp": make_odp(folder / "sample.odp"),
    }
    files["doc"] = lo_convert(files["docx"], "doc:MS Word 97", "doc", folder)
    files["odt"] = lo_convert(files["docx"], "odt", "odt", folder)
    files["rtf"] = lo_convert(files["docx"], "rtf", "rtf", folder)
    files["xls"] = lo_convert(files["xlsx"], "xls:MS Excel 97", "xls", folder)
    files["ods"] = lo_convert(files["xlsx"], "ods", "ods", folder)
    files["pptx"] = lo_convert(files["odp"], "pptx:Impress MS PowerPoint 2007 XML", "pptx", folder)
    files["ppt"] = lo_convert(files["odp"], "ppt:MS PowerPoint 97", "ppt", folder)
    return files


def test_samples_detected_as_expected(samples):
    from konverter.core.detect import detect

    for fmt, path in samples.items():
        detected = detect(path)
        assert detected.fmt == fmt, (fmt, detected)


# --- Конвертации ----------------------------------------------------------------------

CONVERTERS = {
    Category.DOCUMENT: DocumentConverter(),
    Category.SPREADSHEET: SpreadsheetConverter(),
    Category.PRESENTATION: PresentationConverter(),
}


def _convert(converter, fmt: str, source: Path, tid: str, out: Path, cancel=None, preset=Preset.SMALL):
    target = next(t for t in converter.targets(fmt) if t.id == tid)
    assert converter.module_for(tid, fmt) == "libreoffice"
    if target.extension:
        dst = naming.make_temp_file(out, target.extension)
    else:
        dst = naming.make_temp_dir(out)
    values: list[float] = []
    converter.convert([source], dst, target, preset, values.append, cancel or CancelToken(),
                      formats=[fmt])
    assert values and values[-1] == 1.0
    assert all(0.0 <= v <= 1.0 for v in values)
    return dst


def _check_pdf(path: Path, pages_min: int = 1) -> None:
    import pikepdf

    with pikepdf.open(path) as pdf:
        assert len(pdf.pages) >= pages_min
        assert "/Info" not in pdf.trailer
        assert "/Metadata" not in pdf.Root
    assert SECRET_AUTHOR.encode("utf-16-be") not in path.read_bytes()


PDF_PAIRS = [
    (Category.DOCUMENT, "docx"), (Category.DOCUMENT, "doc"), (Category.DOCUMENT, "odt"),
    (Category.DOCUMENT, "rtf"), (Category.SPREADSHEET, "xlsx"), (Category.SPREADSHEET, "xls"),
    (Category.SPREADSHEET, "ods"), (Category.PRESENTATION, "pptx"),
    (Category.PRESENTATION, "ppt"), (Category.PRESENTATION, "odp"),
]


@pytest.mark.parametrize("category,fmt", PDF_PAIRS)
def test_to_pdf(samples, tmp_path, category, fmt):
    dst = _convert(CONVERTERS[category], fmt, samples[fmt], "pdf", tmp_path)
    _check_pdf(dst, 2 if category == Category.PRESENTATION else 1)


@pytest.mark.parametrize("fmt", ["pptx", "ppt", "odp"])
def test_slides_png(samples, tmp_path, fmt):
    from PIL import Image

    dst = _convert(CONVERTERS[Category.PRESENTATION], fmt, samples[fmt], "slides_png", tmp_path)
    pngs = sorted(dst.iterdir())
    assert [p.suffix for p in pngs] == [".png", ".png"]
    with Image.open(pngs[0]) as image:
        assert image.width > image.height > 100


@pytest.mark.parametrize("fmt", ["doc", "odt", "rtf"])
def test_documents_to_docx(samples, tmp_path, fmt):
    import docx

    dst = _convert(CONVERTERS[Category.DOCUMENT], fmt, samples[fmt], "docx", tmp_path)
    document = docx.Document(str(dst))
    assert "ёлка" in "\n".join(p.text for p in document.paragraphs)
    assert document.core_properties.author == ""
    with zipfile.ZipFile(dst) as zf:
        assert SECRET_AUTHOR.encode("utf-8") not in zf.read("docProps/core.xml")


@pytest.mark.parametrize("fmt", ["xls", "ods"])
def test_tables_to_xlsx_and_csv(samples, tmp_path, fmt):
    import openpyxl

    converter = CONVERTERS[Category.SPREADSHEET]
    dst = _convert(converter, fmt, samples[fmt], "xlsx", tmp_path)
    book = openpyxl.load_workbook(dst, read_only=True)
    assert book.sheetnames == ["Лист1", "Второй"]
    book.close()
    with zipfile.ZipFile(dst) as zf:
        core = zf.read("docProps/core.xml")
    assert b"creator" not in core and SECRET_AUTHOR.encode("utf-8") not in core
    csv_path = _convert(converter, fmt, samples[fmt], "csv", tmp_path)
    text = csv_path.read_bytes().decode("utf-8-sig")
    assert "Ёжик;42" in text
    folder = _convert(converter, fmt, samples[fmt], "csv_sheets", tmp_path)
    assert len(list(folder.iterdir())) == 2


def test_cancel_kills_soffice(samples, tmp_path):
    """Отмена во время работы LibreOffice: Cancelled, процессов не остаётся."""
    cancel = CancelToken()
    errors: list[BaseException] = []

    def work() -> None:
        try:
            _convert(CONVERTERS[Category.PRESENTATION], "pptx", samples["pptx"], "pdf", tmp_path,
                     cancel=cancel)
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)

    thread = threading.Thread(target=work)
    thread.start()
    deadline = time.monotonic() + 60
    while not soffice_processes() and time.monotonic() < deadline:
        time.sleep(0.05)
    assert soffice_processes(), "LibreOffice не запустился"
    cancel.cancel()
    started = time.monotonic()
    thread.join(30)
    assert not thread.is_alive()
    assert time.monotonic() - started < 5
    assert errors and isinstance(errors[0], Cancelled)
    assert wait_no_soffice(5) == []


def test_timeout_kills_soffice(samples, tmp_path, monkeypatch):
    monkeypatch.setattr(soffice, "TIMEOUT", 0.5)
    with pytest.raises(ConversionError) as info:
        _convert(CONVERTERS[Category.DOCUMENT], "docx", samples["docx"], "pdf", tmp_path)
    assert str(info.value) == _office.MSG_TIMEOUT
    assert wait_no_soffice(5) == []


def test_damaged_file_gives_error(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"PK\x03\x04" + b"\x00" * 200)
    with pytest.raises(ConversionError):
        _convert(CONVERTERS[Category.DOCUMENT], "docx", bad, "pdf", tmp_path)


def test_user_appdata_untouched(samples, tmp_path):
    """LibreOffice не создаёт профиль в %APPDATA% (всё — в папке программы)."""
    import os

    appdata = Path(os.environ.get("APPDATA", "")) / "LibreOffice"
    existed = appdata.exists()
    _convert(CONVERTERS[Category.DOCUMENT], "odt", samples["odt"], "pdf", tmp_path)
    assert appdata.exists() == existed
