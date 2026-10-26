"""Этап 4: таблицы без LibreOffice — CSV ↔ XLSX (П19) и безопасность.

Проверяется: кодировки и разделители CSV, кириллица в обе стороны, типы
ячеек, защита от внедрения формул (CSV → XLSX: строки, а не формулы;
XLSX → CSV: префикс `'`), несколько листов, XXE и ZIP-бомбы в XLSX, лимиты,
сквозной JobManager.
"""

from __future__ import annotations

import datetime
import os
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

from docs_helpers import (
    LAUGHS,
    bomb_zip,
    file_uri,
    make_inproc,
    many_entries_zip,
    no_children,
    run,
    secret,  # noqa: F401 — фикстура
    target,
    wait_jobs,
)
from konverter.converters import _docjob, tables
from konverter.converters._text import decode_text, sniff_delimiter
from konverter.converters.base import REGISTRY, CancelToken, Category, ConversionError, Preset
from konverter.core.detect import detect
from konverter.core.jobs import JobManager, JobStatus

CONV = tables.SpreadsheetConverter()

ROWS = [
    ["Имя", "Фамилия", "Город", "Сумма"],
    ["Иван", "Петров", "Москва", "1500"],
    ["Мария", "Сидорова", "Санкт-Петербург", "2,5"],
    ["Ёжик", "В тумане", "Тверь", "007"],
]


@pytest.fixture
def inproc(monkeypatch):
    return make_inproc(monkeypatch)


def write_csv(path: Path, rows, encoding: str, delimiter: str) -> Path:
    text = "\r\n".join(delimiter.join(row) for row in rows) + "\r\n"
    path.write_bytes(text.encode(encoding))
    return path


def sheet_values(path: Path) -> list[list]:
    book = openpyxl.load_workbook(path)
    try:
        return [[c.value for c in row] for row in book.worksheets[0].iter_rows()]
    finally:
        book.close()


def csv_text(path: Path) -> str:
    data = path.read_bytes()
    assert data.startswith(b"\xef\xbb\xbf")
    return data[3:].decode("utf-8")


# --- Варианты ---------------------------------------------------------------------------


def test_targets():
    assert isinstance(REGISTRY[Category.SPREADSHEET], tables.SpreadsheetConverter)
    assert [t.id for t in CONV.targets("csv")] == ["xlsx"]
    xlsx = CONV.targets("xlsx")
    # Этап 5: XLSX → PDF, XLS/ODS → XLSX/CSV/PDF — через LibreOffice.
    assert [t.id for t in xlsx] == ["csv", "csv_sheets", "pdf"]
    assert xlsx[0].extension == "csv" and xlsx[1].extension == ""  # папка
    for fmt in ("xls", "ods"):
        assert [t.id for t in CONV.targets(fmt)] == ["xlsx", "csv", "csv_sheets", "pdf"]
    assert CONV.targets("txt") == []
    for t in CONV.targets("csv") + xlsx:
        assert t.is_long and not t.uses_preset and not t.label.startswith("target_")


# --- Кодировки и разделитель ----------------------------------------------------------------


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1251", "utf-16", "utf-16-le"])
def test_decode_text_encodings(encoding):
    text = "Имя;Фамилия\r\nЁжик;Ёлкин\r\nСъешь же ещё этих мягких булок\r\n" * 3
    assert decode_text(text.encode(encoding)) == text


def test_sniff_delimiter():
    assert sniff_delimiter("a,b,c\n1,2,3\n") == ","
    assert sniff_delimiter("a;b;c\n1,5;2,5;3\n") == ";"
    assert sniff_delimiter("a\tb\n1\t2\n") == "\t"
    assert sniff_delimiter('"x;y",b\n"1;2",3\n') == ","
    assert sniff_delimiter("одна колонка\nещё\n") == ","


# --- П19: CSV → XLSX -------------------------------------------------------------------------


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1251", "utf-16"])
@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
def test_p19_csv_to_xlsx(tmp_path, inproc, encoding, delimiter):
    rows = [r[:3] + ([r[3].replace(",", ".")] if delimiter == "," else [r[3]]) for r in ROWS]
    src = write_csv(tmp_path / "данные.csv", rows, encoding, delimiter)
    out = run(CONV, src, tmp_path / "o.xlsx", "csv", "xlsx")
    assert detect(out).fmt == "xlsx"
    values = sheet_values(out)
    assert values[0] == ["Имя", "Фамилия", "Город", "Сумма"]
    assert values[1] == ["Иван", "Петров", "Москва", 1500]
    assert values[2] == ["Мария", "Сидорова", "Санкт-Петербург", 2.5]
    assert values[3] == ["Ёжик", "В тумане", "Тверь", "007"]  # ведущий ноль — текст


def test_csv_quotes_and_newlines(tmp_path, inproc):
    src = tmp_path / "q.csv"
    src.write_bytes('"Поле; с разделителем";"Строка\nвторая";"Кавычка ""внутри"""\r\n1;2;3\r\n'
                    .encode("utf-8"))
    values = sheet_values(run(CONV, src, tmp_path / "o.xlsx", "csv", "xlsx"))
    assert values[0] == ["Поле; с разделителем", "Строка\nвторая", 'Кавычка "внутри"']
    assert values[1] == [1, 2, 3]


def test_formula_injection_csv_to_xlsx(tmp_path, inproc):
    evil = ["=cmd|' /C calc'!A0", "+1+1", "@SUM(1)", "-2+3", "=1+1", "\t=2", "=HYPERLINK(\"x\")"]
    src = write_csv(tmp_path / "evil.csv", [["a", "b", "c", "d", "e", "f", "g"], evil],
                    "utf-8", ";")
    out = run(CONV, src, tmp_path / "o.xlsx", "csv", "xlsx")
    book = openpyxl.load_workbook(out)
    cells = list(book.worksheets[0].iter_rows(min_row=2, max_row=2))[0]
    assert [c.data_type for c in cells] == ["s"] * len(evil)
    assert [c.value for c in cells] == evil
    book.close()
    with zipfile.ZipFile(out) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in sheet and "<f " not in sheet
    assert sheet.count('t="inlineStr"') == 14


def test_csv_limits(tmp_path, inproc, monkeypatch):
    src = write_csv(tmp_path / "big.csv", [[str(i), "x"] for i in range(100)], "utf-8", ";")
    monkeypatch.setitem(_docjob.LIMITS, "max_cells", 50)
    with pytest.raises(ConversionError, match="Слишком большая таблица"):
        run(CONV, src, tmp_path / "o.xlsx", "csv", "xlsx")
    monkeypatch.setitem(_docjob.LIMITS, "max_cells", 5_000_000)
    monkeypatch.setitem(_docjob.LIMITS, "max_rows", 10)
    with pytest.raises(ConversionError, match="Слишком много строк"):
        run(CONV, src, tmp_path / "o2.xlsx", "csv", "xlsx")
    monkeypatch.setitem(_docjob.LIMITS, "max_rows", 1_048_576)
    long_cell = tmp_path / "long.csv"
    long_cell.write_text("a;" + "ж" * 40_000 + "\n", encoding="utf-8")
    with pytest.raises(ConversionError, match="Слишком длинное значение"):
        run(CONV, long_cell, tmp_path / "o3.xlsx", "csv", "xlsx")
    wide = tmp_path / "wide.csv"
    wide.write_text(";".join("x" * 1 for _ in range(16_400)) + "\n", encoding="utf-8")
    with pytest.raises(ConversionError, match="Слишком много столбцов"):
        run(CONV, wide, tmp_path / "o4.xlsx", "csv", "xlsx")


def _local_header(path: Path, name: str) -> dict:
    """Разбор локального заголовка записи (по смещению из центрального каталога)."""
    import struct

    data = path.read_bytes()
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo(name)
    pos = info.header_offset
    assert data[pos:pos + 4] == b"PK\x03\x04"
    (flags, method, crc, csize, usize, name_len, extra_len) = struct.unpack(
        "<HHxxxxIIIHH", data[pos + 6:pos + 30])
    extra = data[pos + 30 + name_len:pos + 30 + name_len + extra_len]
    ids = []
    index = 0
    while index + 4 <= len(extra):
        tag, size = struct.unpack("<HH", extra[index:index + 4])
        ids.append(tag)
        index += 4 + size
    result = {"flags": flags, "crc": crc, "csize": csize, "usize": usize, "extra_ids": ids,
              "info": info}
    if flags & 0x8:  # размеры — в data descriptor после данных
        after = pos + 30 + name_len + extra_len + info.compress_size
        desc = data[after:after + 16]
        if desc[:4] == b"PK\x07\x08":
            desc = desc[4:]
        result["crc"], result["csize"], result["usize"] = struct.unpack("<III", desc[:12])
    return result


def test_xlsx_writer_without_zip64(tmp_path, inproc):
    """Свой XLSX-писатель: у sheet1.xml нет ZIP64-extra, размеры настоящие."""
    src = write_csv(tmp_path / "д.csv", ROWS * 200, "utf-8", ";")
    out = run(CONV, src, tmp_path / "o.xlsx", "csv", "xlsx")
    header = _local_header(out, "xl/worksheets/sheet1.xml")
    info = header["info"]
    assert 0x0001 not in header["extra_ids"]
    assert header["usize"] == info.file_size != 0xFFFFFFFF
    assert header["csize"] == info.compress_size != 0xFFFFFFFF
    assert header["crc"] == info.CRC
    assert 0x0001 not in [t for t in _extra_ids(info.extra)]
    values = sheet_values(out)
    assert values[1] == ["Иван", "Петров", "Москва", 1500] and len(values) == 800


def _extra_ids(extra: bytes) -> list[int]:
    import struct

    ids = []
    index = 0
    while index + 4 <= len(extra):
        tag, size = struct.unpack("<HH", extra[index:index + 4])
        ids.append(tag)
        index += 4 + size
    return ids


def test_xlsx_writer_sheet_size_limit(tmp_path, inproc, monkeypatch):
    """Лист больше предела без ZIP64 — понятная ошибка, а не повреждённый файл."""
    monkeypatch.setattr(tables, "MAX_SHEET_XML_BYTES", 2000)
    src = write_csv(tmp_path / "д.csv", ROWS * 200, "utf-8", ";")
    with pytest.raises(ConversionError, match="Слишком большой результат"):
        run(CONV, src, tmp_path / "o.xlsx", "csv", "xlsx")


# --- П19: XLSX → CSV -------------------------------------------------------------------------


def make_xlsx(path: Path) -> Path:
    book = openpyxl.Workbook()
    first = book.active
    first.title = "Продажи"
    first.append(["Товар", "Цена", "Дата", "Флаг"])
    first.append(["Чай «Цейлон»", 12.5, datetime.date(2024, 1, 31), True])
    first.append(["Кофе; молотый", -3, datetime.datetime(2024, 2, 1, 12, 30), False])
    first.append(["=не формула", "+7 (999)", "@почта", "-минус"])
    first["A4"].data_type = "s"  # текст, начинающийся с «=», а не формула
    first["A5"] = "=1+2"  # формула без кэша значения
    first["A5"].data_type = "f"
    hidden = book.create_sheet("Скрытый")
    hidden.sheet_state = "hidden"
    hidden.append(["тайна"])
    second = book.create_sheet('Лист "2" |3|')
    second.append(["Второй", "лист"])
    second.append([])
    second.append(["после", "пустой"])
    book.create_sheet("Лист <2> <3>")  # после замены символов — то же имя
    book.save(path)
    return path


def test_p19_xlsx_to_csv_first_sheet(tmp_path, inproc):
    src = make_xlsx(tmp_path / "отчёт.xlsx")
    text = csv_text(run(CONV, src, tmp_path / "o.csv", "xlsx", "csv"))
    lines = text.split("\r\n")
    assert lines[0] == "Товар;Цена;Дата;Флаг"
    assert lines[1] == "Чай «Цейлон»;12,5;2024-01-31;TRUE"
    assert lines[2] == '"Кофе; молотый";-3;2024-02-01 12:30:00;FALSE'
    # Защита от формул: текст с =, +, @, - получает префикс «'»; число -3 — нет.
    assert lines[3] == "'=не формула;'+7 (999);'@почта;'-минус"
    assert lines[4] == ""  # формула без кэша — пусто
    assert "тайна" not in text


def test_p19_xlsx_to_csv_every_sheet(tmp_path, inproc):
    src = make_xlsx(tmp_path / "отчёт.xlsx")
    folder = run(CONV, src, tmp_path / "out", "xlsx", "csv_sheets")
    names = sorted(os.listdir(folder))
    assert names == ["Лист _2_ _3_ (2).csv", "Лист _2_ _3_.csv", "Продажи.csv", "Скрытый.csv"]
    assert csv_text(folder / "Лист _2_ _3_.csv") == "Второй;лист\r\n\r\nпосле;пустой\r\n"
    assert csv_text(folder / "Скрытый.csv") == "тайна\r\n"
    assert csv_text(folder / "Лист _2_ _3_ (2).csv") == ""


def test_p19_csv_xlsx_csv_roundtrip(tmp_path, inproc):
    src = write_csv(tmp_path / "исходник.csv", ROWS, "cp1251", ";")
    xlsx = run(CONV, src, tmp_path / "t.xlsx", "csv", "xlsx")
    back = csv_text(run(CONV, xlsx, tmp_path / "back.csv", "xlsx", "csv"))
    assert back == "\r\n".join(";".join(r) for r in ROWS) + "\r\n"


def test_format_value_and_names():
    assert tables.format_value(0.1) == "0,1"
    assert tables.format_value(3.0) == "3"
    assert tables.format_value(None) == ""
    assert tables.format_value("\tтаб") == "'\tтаб"
    assert tables.sheet_file_name("CON", 0, set()) == "_CON.csv"
    assert tables.sheet_file_name("  ", 2, set()) == "Лист3.csv"
    assert tables.column_letters(0) == "A" and tables.column_letters(27) == "AB"
    assert tables.column_letters(16383) == "XFD"


@pytest.mark.parametrize("value", [
    " =1+1", "　=1+1", " +1", "\n=1", "\r\n@SUM(1)", "  -2+3",
    "＝1+1", "＋1", "－1", "＠SUM(1)", "　＝cmd", "|calc", " |calc",
    "\t=2", "\r=2", "=1", "+1", "-x", "@x",
])
def test_csv_formula_prefix_extended(value):
    assert tables.format_value(value) == "'" + value


@pytest.mark.parametrize("value", ["текст", "a=1", "1-2", "", " текст"])
def test_csv_formula_prefix_not_needed(value):
    assert tables.format_value(value) == value


def test_csv_formula_numbers_untouched():
    assert tables.format_value(-3) == "-3"
    assert tables.format_value(-2.5) == "-2,5"


@pytest.mark.parametrize("name", [
    "COM¹", "COM²", "COM³", "LPT¹", "LPT²", "LPT³", "com¹", "lpt³.txt", "Com².csv",
    "con", "Nul.x", "com1", "lpt9.dat", "COM0", "lpt0", "Lpt0.csv", "com0.x",
])
def test_sheet_file_name_reserved(name):
    assert tables.sheet_file_name(name, 0, set()) == "_" + name + ".csv"


def test_sheet_file_name_not_reserved():
    assert tables.sheet_file_name("COM10", 0, set()) == "COM10.csv"
    assert tables.sheet_file_name("Компания", 0, set()) == "Компания.csv"


# --- Безопасность XLSX -----------------------------------------------------------------------


def _xlsx_members(tmp_path: Path) -> dict[str, bytes]:
    src = make_xlsx(tmp_path / "base.xlsx")
    with zipfile.ZipFile(src) as archive:
        return {n: archive.read(n) for n in archive.namelist()}


def _write_members(path: Path, members: dict[str, bytes]) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return path


def test_xxe_in_xlsx(tmp_path, inproc, secret):
    members = _xlsx_members(tmp_path)
    sheet = members["xl/worksheets/sheet1.xml"].decode("utf-8")
    rest = sheet.split("?>", 1)[1] if sheet.startswith("<?xml") else sheet
    decl = "<?xml version=\"1.0\" encoding=\"UTF-8\""
    doctype = f'<!DOCTYPE worksheet [<!ENTITY xxe SYSTEM "{file_uri(secret)}">]>'
    members["xl/worksheets/sheet1.xml"] = (decl + "?>" + doctype + rest).encode("utf-8")
    src = _write_members(tmp_path / "xxe.xlsx", members)
    out = tmp_path / "o.csv"
    with pytest.raises(ConversionError, match="сущност"):
        run(CONV, src, out, "xlsx", "csv")
    assert "СЕКРЕТ" not in out.read_bytes().decode("utf-8", errors="ignore")
    members["xl/worksheets/sheet1.xml"] = (decl + "?>" + LAUGHS.replace("lolz", "worksheet")
                                           + rest).encode("utf-8")
    src2 = _write_members(tmp_path / "lol.xlsx", members)
    t0 = time.monotonic()
    with pytest.raises(ConversionError, match="сущност"):
        run(CONV, src2, tmp_path / "o2.csv", "xlsx", "csv")
    assert time.monotonic() - t0 < 5
    # DOCTYPE в UTF-16 тоже замечается.
    members["xl/worksheets/sheet1.xml"] = (decl.replace("UTF-8", "UTF-16") + "?>" + doctype
                                           + rest).encode("utf-16")
    src3 = _write_members(tmp_path / "u16.xlsx", members)
    with pytest.raises(ConversionError, match="сущност"):
        run(CONV, src3, tmp_path / "o3.csv", "xlsx", "csv")


def test_zip_bombs_in_xlsx(tmp_path, inproc):
    members = _xlsx_members(tmp_path)
    members.pop("xl/sharedStrings.xml", None)
    bomb = bomb_zip(tmp_path / "bomb.xlsx", members=dict(members), bomb_name="xl/sharedStrings.xml")
    t0 = time.monotonic()
    with pytest.raises(ConversionError, match="ZIP-бомб"):
        run(CONV, bomb, tmp_path / "o.csv", "xlsx", "csv")
    assert time.monotonic() - t0 < 10
    many = many_entries_zip(tmp_path / "many.xlsx", first=members)
    with pytest.raises(ConversionError, match="Слишком много частей"):
        run(CONV, many, tmp_path / "o2.csv", "xlsx", "csv")


def test_huge_dimension_sparse_sheet(tmp_path, inproc, monkeypatch):
    """Ячейка в XFD1048576: «пустые» ячейки до неё тоже считаются — быстрый отказ."""
    book = openpyxl.Workbook()
    book.active["A1"] = "a"
    book.active["XFD1048576"] = "z"
    src = tmp_path / "sparse.xlsx"
    book.save(src)
    monkeypatch.setitem(_docjob.LIMITS, "max_cells", 100_000)
    t0 = time.monotonic()
    with pytest.raises(ConversionError, match="Слишком"):
        run(CONV, src, tmp_path / "o.csv", "xlsx", "csv")
    assert time.monotonic() - t0 < 20


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
@pytest.mark.parametrize("payload", [b"PK\x03\x04" + b"\x00" * 64, b"\xff\xfe\x00\x00" * 8,
                                     os.urandom(1024)])
def test_garbage(tmp_path, inproc, fmt, payload):
    src = tmp_path / f"x.{fmt}"
    src.write_bytes(payload)
    tid = "xlsx" if fmt == "csv" else "csv"
    if fmt == "csv":
        # Мусор как CSV — это просто текст: результат есть или понятная ошибка.
        try:
            run(CONV, src, tmp_path / "o.xlsx", fmt, tid)
        except ConversionError:
            pass
        return
    with pytest.raises(ConversionError):
        run(CONV, src, tmp_path / "o.csv", fmt, tid)


# --- JobManager, настоящий ребёнок ------------------------------------------------------------


def test_job_manager_end_to_end(tmp_path):
    folder = tmp_path / "таблицы"
    folder.mkdir()
    src = write_csv(folder / "список.csv", ROWS, "cp1251", ";")
    manager = JobManager(max_workers=1)
    try:
        jobs = [manager.submit([src], CONV, target(CONV, "csv", "xlsx"), Preset.MEDIUM, folder,
                               formats=["csv"])]
        wait_jobs(jobs)
        xlsx = jobs[0].result_path
        jobs.append(manager.submit([xlsx], CONV, target(CONV, "xlsx", "csv_sheets"), Preset.MEDIUM,
                                   folder, formats=["xlsx"]))
        wait_jobs(jobs)
    finally:
        manager.shutdown(timeout=10)
    assert [j.status for j in jobs] == [JobStatus.DONE] * 2, [j.error_message for j in jobs]
    assert jobs[0].result_path == folder / "список.xlsx"
    assert sheet_values(jobs[0].result_path)[3] == ["Ёжик", "В тумане", "Тверь", "007"]
    assert jobs[1].result_path == folder / "список"
    assert os.listdir(jobs[1].result_path) == ["Лист1.csv"]
    assert no_children()


def test_real_child_formula_injection(tmp_path):
    src = write_csv(tmp_path / "e.csv", [["=cmd|' /C calc'!A0", "+1+1"]], "utf-8", ",")
    out = tmp_path / "o.xlsx"
    out.touch()
    CONV.convert([src], out, target(CONV, "csv", "xlsx"), Preset.MEDIUM, lambda _p: None,
                 CancelToken(), formats=["csv"])
    with zipfile.ZipFile(out) as archive:
        assert "<f>" not in archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert no_children()
