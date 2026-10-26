"""Конвертер документов: TXT, HTML, MD, DOCX своим кодом (этап 4); DOC, ODT, RTF и DOCX → PDF — через LibreOffice (этап 5).

Варианты (вход — коды detect), в каноническом порядке DOCX, PDF, HTML, MD, TXT:
* txt, md, html → DOCX, PDF, HTML, MD, TXT (кроме «в тот же формат»);
* docx → PDF (LibreOffice), HTML, MD, TXT;
* doc, odt, rtf → DOCX, PDF (оба через LibreOffice).
Варианты через LibreOffice (`_office`, очередь "office") показываются всегда;
если модуля нет, интерфейс спрашивает разрешение на скачивание (Ф15), а сама
задача без модуля завершается понятной ошибкой.
Все варианты `is_long=True` (дочерний процесс, прогресс, «Отмена»);
пресет ни на что не влияет (`uses_preset=False`).

Чтение (всё — в модель `_docmodel`, она же санитайзер):
* TXT — кодировка по `_text.decode_text` (UTF-8, UTF-8 BOM, UTF-16, cp1251 …);
  абзацы — по пустым строкам, переводы строк внутри абзаца сохраняются;
  если пустых строк нет или почти нет (< 2 % от непустых) — каждая строка
  отдельный абзац (русские TXT-книги так и устроены).
* MD — та же кодировка, Markdown → HTML (Python-Markdown: tables,
  fenced_code, sane_lists), дальше — как HTML: сырой HTML внутри Markdown
  проходит тот же белый список.
* HTML — кодировка: BOM → `<meta charset>` в первых 4 КБ → `decode_text`;
  разбор lxml.html; заголовок страницы `<title>` — название документа.
  Картинки — только `data:image/…;base64` (проверка Pillow, лимиты);
  `http://…`, относительные и локальные пути НЕ загружаются (пропускаются).
* DOCX — после проверки архива (`_safezip.SafeZip.scan`: распаковка всех
  частей «вхолостую» с лимитами + запрет DOCTYPE в XML-частях) — python-docx:
  абзацы и таблицы по порядку, заголовки (стили Heading N / Заголовок N /
  Title), списки (абзацы с нумерацией; маркированный или нумерованный — по
  numFmt), жирный/курсив/моноширинный (Consolas/Courier), ссылки,
  картинки (встроенные; внешние по ссылке — нет), название и автор из
  свойств документа.

Запись:
* HTML — самостоятельная страница UTF-8: CSP `default-src 'none'` (никаких
  скриптов и сети даже если бы что-то просочилось), свой CSS, картинки —
  `data:` URI (суммарно ≤ `max_html_embed_bytes`, остальные пропускаются).
* MD — свой генератор: спецсимволы экранированы, `<` → `&lt;` (сырого HTML в
  результате нет), картинки опускаются с пометкой «[изображение]».
* TXT — UTF-8 без BOM, переводы строк Windows.
* DOCX — python-docx (шаблон по умолчанию самой библиотеки), свойства
  документа очищены (`_reset_properties`): название и язык из источника,
  автор пустой, даты — текущее время UTC, app.xml — «Peregon», без эскиза.
* PDF — `_pdfout` (txt/md/html).
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from konverter.strings import tr

from . import _docjob, _office
from .base import CancelToken, Category, Preset, ProgressCallback, Target

TEXT_INPUTS = frozenset({"txt", "md", "html"})
INPUTS = TEXT_INPUTS | {"docx"}
# Входы и варианты через LibreOffice.
OFFICE_INPUTS = frozenset({"doc", "odt", "rtf"})
_OFFICE_TARGETS = {"docx": ("pdf",), "doc": ("docx", "pdf"), "odt": ("docx", "pdf"),
                   "rtf": ("docx", "pdf")}
# id → (что просит LibreOffice, дообработка `_office`)
_OFFICE_PLAN = {"pdf": ("pdf", "pdf"), "docx": ("docx", "ooxml")}
# (id, расширение)
_OUTPUTS = (
    ("docx", "docx"),
    ("pdf", "pdf"),
    ("html", "html"),
    ("md", "md"),
    ("txt", "txt"),
)
TARGET_ORDER = tuple(item[0] for item in _OUTPUTS)


class DocumentConverter:
    """Конвертер документов (см. докстринг модуля)."""

    category = Category.DOCUMENT
    target_order = TARGET_ORDER
    queue = "docs"

    def targets(self, fmt: str) -> list[Target]:
        fmt = (fmt or "").lower()
        if fmt not in INPUTS and fmt not in OFFICE_INPUTS:
            return []
        office = _OFFICE_TARGETS.get(fmt, ())
        return [
            Target(id=tid, label=tr(f"target_doc_{tid}"), extension=ext, uses_preset=False,
                   is_long=True)
            for tid, ext in _OUTPUTS
            if tid != fmt and (fmt in INPUTS or tid in office)
        ]

    def module_for(self, target_id: str, fmt: str) -> str | None:
        """"libreoffice", если вариант для этого формата делается через LibreOffice."""
        return _office.MODULE if target_id in _OFFICE_TARGETS.get((fmt or "").lower(), ()) else None

    def queue_for(self, target_id: str, formats: list[str]) -> str | None:
        if any(self.module_for(target_id, fmt) for fmt in formats):
            return _office.QUEUE
        return self.queue

    def _base_targets(self, fmt: str) -> list[Target]:
        """Варианты своим кодом (без LibreOffice)."""
        return [t for t in self.targets(fmt) if not self.module_for(t.id, fmt)]

    def convert(
        self,
        sources: list[Path],
        dst: Path,
        target: Target,
        preset: Preset,
        progress: ProgressCallback,
        cancel: CancelToken,
        *,
        formats: list[str],
    ) -> None:
        fmt = str(formats[0] if formats else "").lower()
        if len(formats) == 1 and self.module_for(target.id, fmt):
            kind, post = _OFFICE_PLAN[target.id]
            _office.convert(
                sources, dst, target, preset, progress, cancel, formats, kind=kind, post=post
            )
            return
        _docjob.run_conversion(
            "document", self._base_targets, sources, dst, target, preset, progress, cancel, formats,
            text_formats=TEXT_INPUTS,
        )


# =============================================================================
# Дочерний процесс
# =============================================================================


def child_convert(work) -> None:
    from ._docmodel import doc_to_html, doc_to_markdown, doc_to_text
    from ._isolated import TaskError

    args = work.args
    fmt = args["fmt"]
    target = args["target"]
    src = args["src"]
    if fmt == "docx":
        doc = read_docx(work, src)
    elif fmt in TEXT_INPUTS:
        data = _docjob.read_input(src, _docjob.MAX_TEXT_INPUT_BYTES)
        if fmt == "txt":
            doc = read_txt(work, data)
        elif fmt == "md":
            doc = read_md(work, data)
        else:
            doc = read_html(work, data)
    else:
        raise TaskError("damaged")
    if not doc.chapters or not any(ch.blocks for ch in doc.chapters):
        raise TaskError("empty")
    title_fallback = str(args.get("stem") or "")
    work.tick(0.4)
    if target == "html":
        if not doc.title:
            doc.title = title_fallback
        embed = work.lim("max_html_embed_bytes", 64 * 2**20)
        data = doc_to_html(doc, work, embed_limit=embed).encode("utf-8")
    elif target == "md":
        data = doc_to_markdown(doc, work).replace("\n", "\r\n").encode("utf-8")
    elif target == "txt":
        data = doc_to_text(doc, work, header=False).replace("\n", "\r\n").encode("utf-8")
    elif target == "docx":
        data = write_docx(doc, work)
    elif target == "pdf":
        from ._pdfout import doc_to_pdf

        if not doc.title:
            doc.title = title_fallback
        data = doc_to_pdf(doc, work, book=False, preset=str(args.get("preset", "medium")))
    else:
        raise ValueError("unknown target")
    _docjob.write_bytes(args["dst"], data, work.limits)


# --- Чтение текстовых форматов -------------------------------------------------


def read_txt(work, data: bytes):
    from ._docmodel import Chapter, Doc, Para, Span, clean_text
    from ._text import decode_text

    text = clean_text(decode_text(data))
    work.add_text(len(text))
    all_lines = text.split("\n")
    flags = [bool(line.strip()) for line in all_lines]
    filled = sum(flags)
    # Пустые строки считаются только между текстом (не в начале и не в конце).
    first = flags.index(True) if filled else 0
    last = len(flags) - flags[::-1].index(True) if filled else 0
    empty = (last - first) - filled
    # Пустых строк нет или почти нет (меньше 2 % от непустых: разделители
    # глав в книге) — абзац = строка (так устроены русские TXT-книги).
    # Иначе абзацы — по пустым строкам, переводы строк внутри сохраняются.
    by_line = empty * LINE_PARAGRAPHS_RATIO < filled
    blocks = []
    lines: list[str] = []

    def flush() -> None:
        if lines:
            work.add_block()
            blocks.append(Para([Span("\n".join(line.rstrip() for line in lines))]))
            lines.clear()

    for line in all_lines:
        if line.strip():
            lines.append(line.replace("\t", "    "))
            if by_line:
                flush()
        else:
            flush()
    flush()
    return Doc(chapters=[Chapter("", blocks)] if blocks else [])


# Абзац = строка, если пустых строк меньше, чем непустых / LINE_PARAGRAPHS_RATIO.
LINE_PARAGRAPHS_RATIO = 50


def read_md(work, data: bytes):
    import markdown

    from ._docmodel import Heading, clean_text, spans_text
    from ._text import decode_text

    text = clean_text(decode_text(data))
    work.add_text(len(text))
    work.tick(0.05)
    html = markdown.markdown(
        text, extensions=["tables", "fenced_code", "sane_lists"], output_format="html",
    )
    work.chars = 0  # текст посчитается заново при разборе HTML
    work.tick(0.2)
    doc = _html_doc(work, "<html><body>" + html + "</body></html>")
    if not doc.title:
        for chapter in doc.chapters:
            for block in chapter.blocks:
                if isinstance(block, Heading):
                    doc.title = " ".join(spans_text(block.spans).split())[:500]
                    break
            break
    return doc


_META_CHARSET = re.compile(rb"""<meta[^>]+charset\s*=\s*["']?\s*([A-Za-z0-9_\-:.]+)""", re.I)


def read_html(work, data: bytes):
    from ._text import decode_text

    text = None
    if not data.startswith((b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff")):
        match = _META_CHARSET.search(data[:4096])
        if match:
            try:
                text = data.decode(match.group(1).decode("ascii"), errors="replace")
            except LookupError:
                text = None
    if text is None:
        text = decode_text(data)
    return _html_doc(work, text)


def _html_doc(work, text: str):
    from ._docmodel import Chapter, Doc, HtmlReader, Pictures, data_uri_bytes, html_title
    from ._safexml import parse_html

    root = parse_html(text)
    doc = Doc()
    doc.title = html_title(root)
    lang = root.get("lang") if hasattr(root, "get") else None
    if lang:
        doc.language = lang.strip()[:35]
    pictures = Pictures(work, doc.images)

    def image(src: str) -> str | None:
        if not src.lower().startswith("data:"):
            return None  # внешние и локальные картинки не загружаются
        raw = data_uri_bytes(src, work.max_image_bytes)
        return pictures.add(raw) if raw else None

    blocks = HtmlReader(work, image).blocks(root)
    if blocks:
        doc.chapters.append(Chapter("", blocks))
    return doc


# --- DOCX: чтение --------------------------------------------------------------

_HEADING_STYLE = re.compile(r"^(?:heading|заголовок)\s*(\d)", re.I)
_MONO_FONTS = ("consolas", "courier", "courier new", "lucida console", "menlo", "monaco")


def read_docx(work, path: str):
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.hyperlink import Hyperlink
    from docx.text.paragraph import Paragraph

    from ._docmodel import (
        Chapter, Doc, Heading, Img, ListBlock, Para, Pictures, Span, Table, clean_text,
        merge_spans, safe_href)
    from ._isolated import TaskError
    from ._safezip import precheck_ooxml

    precheck_ooxml(path, work.limits, work.tick)
    work.tick(0.1)
    try:
        document = docx.Document(path)
    except TaskError:
        raise
    except Exception:  # noqa: BLE001 — враждебный файл
        raise TaskError("damaged") from None
    doc = Doc()
    try:
        props = document.core_properties
        doc.title = clean_text(props.title or "")[:1000]
        if props.author:
            doc.authors = [clean_text(props.author)[:300]]
        doc.language = clean_text(props.language or "")[:35]
    except Exception:  # noqa: BLE001
        pass
    pictures = Pictures(work, doc.images)
    part = document.part
    numbering = _numbering_formats(document)

    def image_ids(element) -> list[str]:
        ids = []
        for rid in element.xpath(".//a:blip/@r:embed"):
            related = part.related_parts.get(rid)
            blob = getattr(related, "blob", None)
            if blob:
                image_id = pictures.add(blob, key=rid)
                if image_id:
                    ids.append(image_id)
        return ids

    def run_spans(run, href=None) -> tuple[list, list]:
        text = clean_text(run.text or "")
        font = None
        try:
            font = (run.font.name or "").lower()
        except Exception:  # noqa: BLE001
            font = ""
        code = any(font.startswith(m) for m in _MONO_FONTS) if font else False
        spans = [Span(text, bool(run.bold), bool(run.italic), code, href)] if text else []
        return spans, image_ids(run._r)

    def paragraph_content(paragraph) -> tuple[list, list]:
        spans: list = []
        images: list = []
        for item in paragraph.iter_inner_content():
            if isinstance(item, Hyperlink):
                href = None
                try:
                    href = safe_href(item.address) if item.address else None
                except Exception:  # noqa: BLE001
                    href = None
                for run in item.runs:
                    s, i = run_spans(run, href)
                    spans.extend(s)
                    images.extend(i)
            else:
                s, i = run_spans(item)
                spans.extend(s)
                images.extend(i)
        return merge_spans(spans), images

    def list_info(paragraph) -> tuple[int, bool] | None:
        ppr = paragraph._p.pPr
        num_pr = ppr.numPr if ppr is not None else None
        style = ""
        try:
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
        except Exception:  # noqa: BLE001
            style = ""
        if num_pr is None:
            if style.lower().startswith(("list bullet", "list number")):
                return 0, "number" in style.lower()
            return None
        level = 0
        num_id = None
        if num_pr.ilvl is not None and num_pr.ilvl.val is not None:
            level = max(0, min(8, int(num_pr.ilvl.val)))
        if num_pr.numId is not None:
            num_id = num_pr.numId.val
        if num_id == 0:
            return None
        fmt = numbering.get((num_id, level))
        ordered = fmt not in (None, "bullet", "none") if fmt else "number" in style.lower()
        return level, ordered

    blocks: list = []
    list_stack: list[tuple[int, ListBlock]] = []

    def close_lists() -> None:
        list_stack.clear()

    def add_list_item(level: int, ordered: bool, item_blocks: list) -> None:
        while list_stack and list_stack[-1][0] > level:
            list_stack.pop()
        if list_stack and list_stack[-1][0] == level and list_stack[-1][1].ordered != ordered:
            list_stack.pop()  # другой вид списка на том же уровне — новый список
        if not list_stack or list_stack[-1][0] < level:
            new = ListBlock(ordered, [])
            if list_stack and list_stack[-1][1].items:
                list_stack[-1][1].items[-1].append(new)
            else:
                blocks.append(new)
            list_stack.append((level, new))
        list_stack[-1][1].items.append(item_blocks)

    def handle_paragraph(paragraph) -> None:
        spans, images = paragraph_content(paragraph)
        style = ""
        try:
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
        except Exception:  # noqa: BLE001
            style = ""
        heading = _HEADING_STYLE.match(style)
        info = list_info(paragraph) if not heading else None
        if info is not None and spans:
            work.add_text(sum(len(s.text) for s in spans))
            work.add_block()
            add_list_item(info[0], info[1], [Para(spans)])
            for image_id in images:
                blocks.append(Img(image_id))
            return
        close_lists()
        if spans and "".join(s.text for s in spans).strip():
            work.add_text(sum(len(s.text) for s in spans))
            work.add_block()
            if heading or style.lower() == "title":
                level = int(heading.group(1)) if heading else 1
                flat = [Span(s.text.replace("\n", " "), s.bold, s.italic, s.code, s.href)
                        for s in spans]
                blocks.append(Heading(max(1, min(6, level)), flat))
                if style.lower() == "title" and not doc.title:
                    doc.title = " ".join("".join(s.text for s in spans).split())[:500]
            else:
                blocks.append(Para(spans))
        for image_id in images:
            work.add_block()
            blocks.append(Img(image_id))

    def handle_table(table) -> None:
        close_lists()
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                spans: list = []
                for paragraph in cell.paragraphs:
                    content, _images = paragraph_content(paragraph)
                    if content:
                        if spans:
                            spans.append(Span("\n"))
                        spans.extend(content)
                cells.append(merge_spans(spans))
            work.add_cells(len(cells))
            rows.append(cells)
        if rows:
            work.add_block()
            blocks.append(Table(rows, 0))

    def walk(container) -> None:
        for index, child in enumerate(container.iterchildren()):
            if child.tag == qn("w:p"):
                handle_paragraph(Paragraph(child, document._body))
            elif child.tag == qn("w:tbl"):
                handle_table(DocxTable(child, document._body))
            elif child.tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    walk(content)
            if index % 200 == 0:
                work.tick()

    try:
        walk(document.element.body)
    except TaskError:
        raise
    except Exception:  # noqa: BLE001 — нестандартная структура
        raise TaskError("damaged") from None
    if blocks:
        doc.chapters.append(Chapter("", blocks))
    return doc


def _numbering_formats(document) -> dict:
    """(numId, уровень) → numFmt ("bullet", "decimal" …) из numbering.xml."""
    from docx.oxml.ns import qn

    result: dict = {}
    try:
        numbering = document.part.numbering_part.element
    except Exception:  # noqa: BLE001 — нумерации нет
        return result
    try:
        abstract: dict = {}
        for node in numbering.findall(qn("w:abstractNum")):
            levels = {}
            for lvl in node.findall(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                levels[int(lvl.get(qn("w:ilvl"), "0"))] = (
                    fmt.get(qn("w:val")) if fmt is not None else None)
            abstract[node.get(qn("w:abstractNumId"))] = levels
        for num in numbering.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            if ref is None:
                continue
            num_id = int(num.get(qn("w:numId"), "0"))
            for level, fmt in abstract.get(ref.get(qn("w:val")), {}).items():
                result[(num_id, level)] = fmt
    except (ValueError, TypeError):
        return result
    return result


# --- DOCX: запись --------------------------------------------------------------


# docProps/app.xml результата вместо шаблонного python-docx («Microsoft
# Macintosh Word», «Normal.dotm», AppVersion 14 и счётчики пустого документа).
_APP_XML = (
    b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n"
    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
    b'extended-properties"><Application>Peregon</Application></Properties>'
)


def _reset_properties(document, doc) -> None:
    """Свойства DOCX-результата: только название и язык из источника.

    Шаблон python-docx несёт чужие свойства: создан/изменён 2013-12-23, автор
    «python-docx», описание «generated by python-docx», app.xml от Mac Word,
    эскиз пустой страницы. Даты — текущее время UTC; автор и «кем изменён»
    пустые: это документ, а не книга — автор исходника (из свойств DOCX) не
    переносится, чтобы не тащить чужие персональные данные.
    """
    import datetime

    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
    props = document.core_properties
    props.title = doc.title or ""
    props.author = ""
    props.comments = ""
    props.keywords = ""
    props.subject = ""
    props.category = ""
    props.last_modified_by = ""
    props.revision = 1
    props.created = now
    props.modified = now
    if doc.language:
        props.language = doc.language
    package = document.part.package
    for r_id, rel in list(package.rels.items()):
        if rel.reltype == RT.THUMBNAIL:
            package.rels.pop(r_id)  # эскиз шаблона (пустая страница) — убрать
        elif rel.reltype == RT.EXTENDED_PROPERTIES and not rel.is_external:
            rel.target_part._blob = _APP_XML


def write_docx(doc, work) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    from ._docmodel import Heading, Img, ListBlock, Para, Pre, Quote, Rule, Span, Table

    document = docx.Document()
    _reset_properties(document, doc)
    styles = {s.name for s in document.styles}

    def add_hyperlink(paragraph, span) -> None:
        r_id = paragraph.part.relate_to(span.href, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        lines = span.text.split("\n")
        for index, line in enumerate(lines):
            run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            for tag, value in (("w:color", "1F4EB4"), ("w:u", "single")):
                element = OxmlElement(tag)
                element.set(qn("w:val"), value)
                rpr.append(element)
            if span.bold:
                rpr.append(OxmlElement("w:b"))
            if span.italic:
                rpr.append(OxmlElement("w:i"))
            run.append(rpr)
            if index:
                run.append(OxmlElement("w:br"))
            text = OxmlElement("w:t")
            text.set(qn("xml:space"), "preserve")
            text.text = line
            run.append(text)
            link.append(run)
        paragraph._p.append(link)

    def runs(paragraph, spans, *, bold=False, italic=False) -> None:
        for span in spans:
            if not span.text:
                continue
            if span.href:
                add_hyperlink(paragraph, span)
                continue
            run = paragraph.add_run(span.text)
            if span.bold or bold:
                run.bold = True
            if span.italic or italic:
                run.italic = True
            if span.code:
                run.font.name = "Consolas"

    def list_style(ordered: bool, depth: int) -> str | None:
        base = "List Number" if ordered else "List Bullet"
        name = base if depth == 0 else f"{base} {min(3, depth + 1)}"
        return name if name in styles else (base if base in styles else None)

    def add_blocks(blocks, depth: int = 0, quote: bool = False) -> None:
        for block in blocks:
            work.add_block()
            add_block(block, depth, quote)

    def add_block(block, depth: int, quote: bool) -> None:
        if isinstance(block, Heading):
            paragraph = document.add_heading("", level=max(1, min(9, block.level)))
            runs(paragraph, block.spans)
        elif isinstance(block, Para):
            style = "Quote" if quote and "Quote" in styles else None
            paragraph = document.add_paragraph(style=style)
            runs(paragraph, block.spans, italic=quote and style is None)
        elif isinstance(block, ListBlock):
            for number, item in enumerate(block.items, 1):
                first = True
                for sub in item:
                    if first and isinstance(sub, Para):
                        paragraph = document.add_paragraph(style=list_style(block.ordered, depth))
                        runs(paragraph, sub.spans)
                    elif isinstance(sub, ListBlock):
                        add_block(sub, depth + 1, quote)
                    else:
                        add_block(sub, depth, quote)
                    first = False
        elif isinstance(block, Table):
            width = max(len(row) for row in block.rows)
            if width > 63:  # предел столбцов Word
                for row in block.rows:
                    paragraph = document.add_paragraph()
                    flat: list = []
                    for index, cell in enumerate(row):
                        if index:
                            flat.append(Span("\t"))
                        flat.extend(cell)
                    runs(paragraph, flat)
                return
            table = document.add_table(rows=0, cols=width)
            if "Table Grid" in styles:
                table.style = "Table Grid"
            for row_index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for index, cell in enumerate(row):
                    runs(cells[index].paragraphs[0], cell, bold=row_index < block.header)
                if row_index % 200 == 0:
                    work.tick()
        elif isinstance(block, Img):
            picture = doc.images.get(block.id)
            if picture is None:
                return
            width_in = min(6.0, picture.width / 96)
            try:
                document.add_picture(io.BytesIO(picture.data), width=Inches(max(0.2, width_in)))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:  # noqa: BLE001 — картинку python-docx не понял
                return
        elif isinstance(block, Pre):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
        elif isinstance(block, Quote):
            add_blocks(block.blocks, depth, True)
        elif isinstance(block, Rule):
            paragraph = document.add_paragraph("* * *")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    total = max(1, len(doc.chapters))
    for index, chapter in enumerate(doc.chapters):
        if index:
            document.add_page_break()
        add_blocks(chapter.blocks)
        work.tick(0.4 + 0.5 * (index + 1) / total)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
